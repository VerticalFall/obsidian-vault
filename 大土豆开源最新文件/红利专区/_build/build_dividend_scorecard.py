from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.pagebreak import Break
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "红利可持续性评分表.xlsx"

# ---------- style tokens ----------
NAVY = "1F3864"
BLUE = "2F5597"
MID_BLUE = "D9E2F3"
LIGHT_BLUE = "DDEBF7"
LIGHT_GRAY = "F2F2F2"
LIGHT_YELLOW = "FFF2CC"
LIGHT_RED = "FCE4D6"
LIGHT_GREEN = "E2EFDA"
WHITE = "FFFFFF"
BORDER = "B4C6E7"
TEXT = "333333"
MUTED = "5B6B7C"
RED = "C00000"
GREEN = "2E7D32"

thin = Side(style="thin", color=BORDER)
BORDER_ALL = Border(left=thin, right=thin, top=thin, bottom=thin)

FONT = "微软雅黑"

def font(size=10, bold=False, color=TEXT, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)

def fill(color):
    return PatternFill("solid", fgColor=color)

def align(h="center", v="center", wrap=True):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def cell(ws, r, c, value=None, kind="body", fmt=None, bold=False,
         h="center", wrap=True, size=10, color=None):
    x = ws.cell(row=r, column=c, value=value)
    x.border = BORDER_ALL
    x.alignment = align(h, "center", wrap)
    if kind == "header":
        x.fill = fill(NAVY)
        x.font = font(size, True, WHITE)
    elif kind == "section":
        x.fill = fill(MID_BLUE)
        x.font = font(size, True, NAVY)
    elif kind == "input":
        x.fill = fill(LIGHT_BLUE)
        x.font = font(size, bold, color or BLUE)
    elif kind == "formula":
        x.fill = fill(LIGHT_GRAY)
        x.font = font(size, bold, color or TEXT)
    elif kind == "assumption":
        x.fill = fill(LIGHT_YELLOW)
        x.font = font(size, bold, color or TEXT)
    elif kind == "risk":
        x.fill = fill(LIGHT_RED)
        x.font = font(size, bold, color or RED)
    elif kind == "note":
        x.fill = fill(WHITE)
        x.font = font(size, bold, color or MUTED, italic=True)
    else:
        x.fill = fill(WHITE)
        x.font = font(size, bold, color or TEXT)
    if fmt:
        x.number_format = fmt
    return x

def title(ws, row, text, span, sub=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row=row, column=1, value=text)
    c.fill = fill(NAVY)
    c.font = font(16, True, WHITE)
    c.alignment = align("left", "center", False)
    ws.row_dimensions[row].height = 32
    if sub:
        ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=span)
        c2 = ws.cell(row=row + 1, column=1, value=sub)
        c2.fill = fill(LIGHT_BLUE)
        c2.font = font(10, False, MUTED)
        c2.alignment = align("left", "center", True)
        ws.row_dimensions[row + 1].height = 26

def split_title(ws, row, text, span, sub=None):
    """宽表按 2 页宽打印时，标题只放在 A:C，避免第二页被截断。"""
    for c in range(1, span + 1):
        ws.cell(row=row, column=c).fill = fill(NAVY)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    c = ws.cell(row=row, column=1, value=text)
    c.fill = fill(NAVY)
    c.font = font(15, True, WHITE)
    c.alignment = align("left", "center", False)
    ws.row_dimensions[row].height = 32
    if sub:
        for cc in range(1, span + 1):
            ws.cell(row=row + 1, column=cc).fill = fill(LIGHT_BLUE)
        ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=3)
        c2 = ws.cell(row=row + 1, column=1, value=sub)
        c2.fill = fill(LIGHT_BLUE)
        c2.font = font(9, False, MUTED)
        c2.alignment = align("left", "center", True)
        ws.row_dimensions[row + 1].height = 30


def header_row(ws, row, headers, height=38):
    for i, h in enumerate(headers, 1):
        cell(ws, row, i, h, "header", size=10)
    ws.row_dimensions[row].height = height

def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w

def merge_note(ws, row, text, span, height=26):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row=row, column=1, value=text)
    c.fill = fill(LIGHT_YELLOW)
    c.font = font(10, False, TEXT)
    c.alignment = align("left", "center", True)
    c.border = BORDER_ALL
    ws.row_dimensions[row].height = height

def page(ws, paper=None, pages_wide=1, title_cols=None, pages_high=0):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = pages_wide
    ws.page_setup.fitToHeight = pages_high
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    if paper:
        ws.page_setup.paperSize = paper
    if title_cols:
        ws.print_title_cols = title_cols

wb = Workbook()

# ============================================================
# 说明
# ============================================================
ws = wb.active
ws.title = "说明"
ws.sheet_properties.tabColor = "1F3864"
widths(ws, {"A": 20, "B": 28, "C": 36, "D": 36, "E": 22, "F": 24})
title(ws, 1, "红利可持续性评分表 · 使用说明", 6, "v1.0.0 · 2026-10-30 · 研究分级工具，不构成投资建议")

cell(ws, 3, 1, "定位", "section", h="left")
ws.merge_cells("A4:F4")
c = ws.cell(row=4, column=1, value="把红利从“高股息率”一个标签，拆成三条轨道：一般红利/类债、金融红利、质量防御。质量分和估值分分开，硬约束优先于加权分。")
c.fill = fill(WHITE); c.font = font(10); c.alignment = align("left", "center", True); c.border = BORDER_ALL
ws.row_dimensions[4].height = 38

cell(ws, 6, 1, "三条轨道", "section", h="left")
ws.merge_cells("A6:F6")
header_row(ws, 7, ["轨道", "适用", "核心问题", "评分方法", "权威来源", "备注"])
tracks = [
    ("一般红利 / 类债", "非金融、稳定分红公司", "FCF 能不能覆盖股息？", "六维质量分 + 硬约束", "本框架", "本表主评分区"),
    ("金融红利", "银行、保险", "资本、资产质量、股息能否持续？", "银行 skill 四维评分 + DDM", "银行估值skills", "不套用 FCF 评分"),
    ("质量防御", "Visa、Mastercard、伯克希尔等", "生意质量能否穿越周期？", "五维质量分 + 估值分", "本框架", "与红利分开排序"),
]
for r, row in enumerate(tracks, 8):
    for c_, v in enumerate(row, 1):
        cell(ws, r, c_, v, "body", h="left" if c_ != 1 else "center", size=9)
    ws.row_dimensions[r].height = 30

cell(ws, 12, 1, "评级", "section", h="left")
ws.merge_cells("A12:F12")
header_row(ws, 13, ["等级", "最终质量分", "含义", "研究动作"])
cell(ws, 13, 5, None, "header")
cell(ws, 13, 6, None, "header")
ws.merge_cells("D13:F13")
grades = [
    ("A", "≥80", "红利质量强", "进入核心研究池"),
    ("B", "65–79", "合格", "持续跟踪"),
    ("C", "50–64", "观察", "需要催化剂或改善证据"),
    ("D", "<50", "不满足红利标准", "不进入红利池"),
    ("NA", "数据不足 / 负 FCF / 不适用", "不强行评分", "补数据或换方法"),
]
for r, row in enumerate(grades, 14):
    for c_, v in enumerate(row, 1):
        cell(ws, r, c_, v, "body", h="left" if c_ in (3, 4) else "center", size=9)
    for c_ in (5, 6):
        cell(ws, r, c_, None, "body")
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)

cell(ws, 20, 1, "硬约束（优先于加权分）", "section", h="left")
ws.merge_cells("A20:F20")
notes = [
    "FCF 派息率连续 2 年 >100% → 最终质量分上限 40。",
    "净负债 / EBITDA >4 倍 → 上限 30。",
    "近 5 年砍息 → 上限 50。",
    "5 年归母净利最大回撤 >70% → 上限 50。",
    "归母与普通股东利润差异 >20%（少数股东/治理侵蚀）→ 上限 60。",
    "硬约束触发后：最终质量分 = MIN(加权质量分, 上限)，不因为其他维度优秀而抵消。",
]
for i, n in enumerate(notes, 21):
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=6)
    c = ws.cell(row=i, column=1, value="· " + n)
    c.fill = fill(LIGHT_RED); c.font = font(9, False, TEXT); c.alignment = align("left", "center", True); c.border = BORDER_ALL
    ws.row_dimensions[i].height = 22

cell(ws, 28, 1, "总回报公式", "section", h="left")
ws.merge_cells("A28:F28")
formula_notes = [
    "税后股息率 = 股息率 × (1 − 适用税率)",
    "总回报 = (1 + 税后股息率) × (1 + 价格回报) × (1 + 汇率变动) − 1",
    "超额收益 = 总回报 − 本地 10 年期国债收益率",
]
for i, n in enumerate(formula_notes, 29):
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=6)
    c = ws.cell(row=i, column=1, value=n)
    c.fill = fill(LIGHT_BLUE); c.font = font(10, False, TEXT); c.alignment = align("left", "center", True); c.border = BORDER_ALL
    ws.row_dimensions[i].height = 22

cell(ws, 33, 1, "使用步骤", "section", h="left")
ws.merge_cells("A33:F33")
steps = [
    "1. 在「评分总表」选轨道：一般红利 / 金融红利 / 质量防御。",
    "2. 一般红利：在「分红安全」填当前价、DPS、股本、FCF、净负债/EBITDA、capex/折旧、5 年净利回撤、砍息次数、治理分。",
    "3. 金融红利：从银行 skill 的四维评分粘贴到「金融红利」，不要用 FCF 评分。",
    "4. 质量防御：在「质量防御」填 ROIC/WACC、capex/营收、FCF 波动、净现金/市值、回购收益率、DCF 基准价。",
    "5. 在「总回报」填税率、价格回报、汇率，得到税后总回报和相对 10Y 的超额。",
    "6. 在「组合监控」填目标权重，检查行业集中、利率敏感度、触发条件。",
    "7. 最后看「检查」页：权重、范围、重复代码、数据完整度是否通过。",
]
for i, s in enumerate(steps, 34):
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=6)
    c = ws.cell(row=i, column=1, value=s)
    c.fill = fill(WHITE); c.font = font(9, False, TEXT); c.alignment = align("left", "center", True); c.border = BORDER_ALL
    ws.row_dimensions[i].height = 24

cell(ws, 42, 1, "版本与免责", "section", h="left")
ws.merge_cells("A42:F42")
ws.merge_cells("A43:F43")
c = ws.cell(row=43, column=1, value="v1.0.0 / 2026-10-30。本表只做研究分级与组合讨论，不构成任何投资建议。税率、汇率、价格回报均为输入假设，请按实际口径填写。")
c.fill = fill(LIGHT_YELLOW); c.font = font(9, False, TEXT); c.alignment = align("left", "center", True); c.border = BORDER_ALL
ws.row_dimensions[43].height = 34

cell(ws, 45, 1, "读表提示", "section", h="left")
ws.merge_cells("A45:F45")
ws.merge_cells("A46:F46")
c = ws.cell(row=46, column=1, value="本工作簿是“模板 + 已知硬数据示例”：空白 = 待填，不按 0 计算；六维未填满时质量分显示“数据不足”，不是 0 分。预填的 7 家一般红利、4 家银行、3 家质量防御只用于演示口径，不代表已完成评分。")
c.fill = fill(WHITE); c.font = font(9); c.alignment = align("left", "center", True); c.border = BORDER_ALL
ws.row_dimensions[46].height = 46

# 让“总回报公式”标题与内容留在同一页（断点在标题行之前）
ws.row_breaks.append(Break(id=27))
page(ws)

# ============================================================
# 评分总表
# ============================================================
ws = wb.create_sheet("评分总表")
ws.sheet_properties.tabColor = "2F5597"
widths(ws, {"A": 5, "B": 16, "C": 13, "D": 8, "E": 15, "F": 15, "G": 9, "H": 9, "I": 9, "J": 7,
            "K": 9, "L": 9, "M": 10, "N": 30, "O": 38, "P": 11, "Q": 34})
title(ws, 1, "评分总表 · 三条轨道总览", 17,
      "质量分与估值分分开；综合分 = 60% 质量 + 40% 估值；等级只代表研究优先级，不代表买卖建议。")
headers = ["序号", "标的", "代码", "市场", "轨道", "行业", "质量分", "估值分", "综合分", "等级",
           "股息率", "总回报", "总回报-10Y", "硬约束 / 口径", "结论", "数据日期", "备注"]
header_row(ws, 4, headers, 42)

rows = [
    ("工商银行", "601398.SH", "A股", "金融红利", "银行", "合理偏低；ROE 低于股权成本", "2026-10-30", "银行类债；四维评分引用银行 skill"),
    ("平安银行", "000001.SZ", "A股", "金融红利", "银行", "合理偏贵；ROE 低于股权成本", "2026-10-30", "DDM 8.13，综合基准 11.03；现价 11.64"),
    ("美的集团", "000333.SZ", "A股", "一般红利", "家电", "现价低于基准 22%，便宜", "2026-10-30", "15 倍 PE；4.9% 股息率；净现金 455 亿"),
    ("贵州茅台", "600519.SH", "A股", "一般红利", "高端白酒", "接近内在价值，合理", "2026-10-30", "20.8 倍 PE；3.8% 股息率"),
    ("长江电力", "600900.SH", "A股", "一般红利", "水电", "合理偏贵；好资产，价格在上沿", "2026-10-30", "22.6 倍 PE；3.1% 股息率；基准 27.59"),
    ("中国神华", "601088.SH", "A股", "一般红利", "煤炭", "昂贵；为周期高位盈利付费", "2026-10-30", "20 倍 PE；3.8% 股息率；基准 36.64"),
    ("可口可乐", "KO", "美股", "一般红利", "必选消费", "护城河仍在，估值略贵", "2026-10-30", "27 倍 PE；股息率待填"),
    ("中国移动", "600941.SH", "A股", "一般红利", "运营商", "待研究", "", "观察池候选；资本开支见顶后分红提升"),
    ("宁沪高速", "600377.SH", "A股", "一般红利", "高速公路", "待研究", "", "观察池候选；类债现金流"),
    ("伯克希尔", "BRK.B", "美股", "质量防御", "综合控股", "合理区间上沿，安全边际不足", "2026-10-30", "非股息型；现金与上市股权提供资产负债表防御"),
    ("Visa", "V", "美股", "质量防御", "支付", "合理偏低，安全边际不足", "2026-10-30", "32.5 倍 PE；0.78% 股息率；DCF 424.63"),
    ("万事达卡", "MA", "美股", "质量防御", "支付", "合理偏低，安全边际不足", "2026-10-30", "33.6 倍 PE；0.50% 股息率；DCF 666.31"),
]

# row mapping to detail sheets
# 分红安全 rows 5-11: 美的, 茅台, 长江, 神华, KO, 中国移动, 宁沪
# 金融红利 rows 5-8: 工行, 平安, 北京, 宁波
# 质量防御 rows 5-7: 伯克希尔, Visa, MA
general_map = {7: 5, 8: 6, 9: 7, 10: 8, 11: 9, 12: 10, 13: 11}
quality_map = {14: 5, 15: 6, 16: 7}

for i, (name, code, market, track, industry, conclusion, data_date, note) in enumerate(rows, 5):
    cell(ws, i, 1, i - 4, "body")
    cell(ws, i, 2, name, "body", h="left", bold=True)
    cell(ws, i, 3, code, "body")
    cell(ws, i, 4, market, "body")
    cell(ws, i, 5, track, "input")
    cell(ws, i, 6, industry, "body")
    # quality score
    if track == "金融红利":
        src = 5 if name == "工商银行" else 6
        cell(ws, i, 7, f"='金融红利'!J{src}", "formula", fmt="0.0")
        cell(ws, i, 14, "银行资本/资产质量约束，见银行 skill", "note", h="left", size=9)
    elif track == "质量防御":
        src = quality_map[i]
        cell(ws, i, 7, f"='质量防御'!R{src}", "formula", fmt="0.0")
        cell(ws, i, 14, "ROIC / 资本强度约束，见质量防御", "note", h="left", size=9)
    else:
        src = general_map[i]
        # 未填满六维时明确显示“数据不足”，而不是留空被误读为 0 分
        cell(ws, i, 7, f'=IF(\'分红安全\'!AD{src}="数据不足","数据不足",\'分红安全\'!AH{src})', "formula", fmt="0.0")
        cell(ws, i, 14, f"='分红安全'!AG{src}", "note", h="left", size=9)
    # valuation score
    cell(ws, i, 8, f"='总回报'!T{i}", "formula", fmt="0.0")
    # composite: 只有质量分和估值分都是数值才计算，文本“数据不足”不参与运算
    cell(ws, i, 9, f'=IF(AND(ISNUMBER(G{i}),ISNUMBER(H{i})),0.6*G{i}+0.4*H{i},"")', "formula", fmt="0.0")
    cell(ws, i, 10, f'=IF(ISNUMBER(I{i}),IF(I{i}>=80,"A",IF(I{i}>=65,"B",IF(I{i}>=50,"C","D"))),"")', "formula", bold=True)
    cell(ws, i, 11, f"='总回报'!H{i}", "formula", fmt="0.00%")
    cell(ws, i, 12, f"='总回报'!P{i}", "formula", fmt="0.00%")
    cell(ws, i, 13, f"='总回报'!Q{i}", "formula", fmt="0.00%")
    cell(ws, i, 14 if track != "一般红利" else 14, ws.cell(row=i, column=14).value, "note", h="left", size=9)
    cell(ws, i, 15, conclusion, "body", h="left", size=9)
    cell(ws, i, 16, data_date, "body", fmt="yyyy-mm-dd")
    cell(ws, i, 17, note, "note", h="left", size=9)
    ws.row_dimensions[i].height = 40

# Overwrite N for general rows with formula after initial loop (keeps style)
for i, src in general_map.items():
    cell(ws, i, 14, f"='分红安全'!AG{src}", "note", h="left", size=9)
for i in (5, 6):
    cell(ws, i, 14, "银行资本/资产质量约束，见银行 skill", "note", h="left", size=9)
for i in (14, 15, 16):
    cell(ws, i, 14, "ROIC / 资本强度约束，见质量防御", "note", h="left", size=9)

# drop unused mapped cell writes; ensure formula in N didn't get duplicated
for i in range(5, 17):
    v = ws.cell(row=i, column=14).value
    if v is None:
        cell(ws, i, 14, "", "note", h="left", size=9)

ws.freeze_panes = "C5"
dv_track = DataValidation(type="list", formula1='"一般红利,一般红利/类债,金融红利,质量防御"', allow_blank=True)
ws.add_data_validation(dv_track); dv_track.add("E5:E16")
for val, color in [("A", LIGHT_GREEN), ("B", MID_BLUE), ("C", LIGHT_YELLOW), ("D", LIGHT_RED)]:
    ws.conditional_formatting.add("J5:J16", FormulaRule(formula=[f'J5="{val}"'], fill=fill(color)))
page(ws, paper=8)

# ============================================================
# 分红安全
# ============================================================
ws = wb.create_sheet("分红安全")
ws.sheet_properties.tabColor = "C77C1A"
widths(ws, {"A": 5, "B": 15, "C": 13, "D": 8, "E": 14, "F": 9, "G": 10, "H": 10, "I": 12, "J": 9,
            "K": 12, "L": 11, "M": 12, "N": 10, "O": 11, "P": 11, "Q": 11, "R": 12, "S": 11, "T": 11,
            "U": 10, "V": 11, "W": 12, "X": 12, "Y": 10, "Z": 12, "AA": 10, "AB": 10, "AC": 11,
            "AD": 10, "AE": 14, "AF": 10, "AG": 34, "AH": 11, "AI": 8, "AJ": 10, "AK": 42})
# 打印按 2 页宽拆分：标题放 A:C，保证第二页完整显示
split_title(ws, 1, "一般红利 · 六维评分", 24,
            "蓝色=输入；灰色=公式；红色=硬约束")
headers = ["序号", "标的", "代码", "市场", "行业", "当前价", "2026E DPS", "股本(亿股)", "2026E 分红总额(亿元)",
           "股息率", "FY2025 FCF(亿元)", "FCF派息率", "3年平均FCF派息率", "FCF覆盖得分",
           "净负债(亿元)", "EBITDA(亿元)", "净负债/EBITDA", "资产负债表得分",
           "资本开支(亿元)", "折旧摊销(亿元)", "capex/折旧", "资本开支得分",
           "5年净利最大回撤", "盈利周期性得分", "近5年砍息次数", "分红连续性得分",
           "5年DPS CAGR", "治理得分", "少数股东占比", "质量分", "连续2年>100%",
           "硬约束上限", "硬约束原因", "最终质量分", "等级", "数据完整度", "备注"]
header_row(ws, 4, headers, 58)

general = [
    ("美的集团", "000333.SZ", "A股", "家电", "股息率约4.9%；FCF/净负债/资本开支待填"),
    ("贵州茅台", "600519.SH", "A股", "高端白酒", "股息率约3.8%；FCF/净负债/资本开支待填"),
    ("长江电力", "600900.SH", "A股", "水电", "净负债2734.9亿元；少数股东119.8亿元；FCF/EBITDA/资本开支待填"),
    ("中国神华", "601088.SH", "A股", "煤炭", "FY2025 FCF 266.59亿元；2026E DPS 1.95元；单年FCF派息率约159%，需核连续性与资本开支口径"),
    ("可口可乐", "KO", "美股", "必选消费", "PE约27倍；股息率/FCF待填"),
    ("中国移动", "600941.SH", "A股", "运营商", "待研究；观察池候选"),
    ("宁沪高速", "600377.SH", "A股", "高速公路", "待研究；观察池候选"),
]
for i, (name, code, market, industry, note) in enumerate(general, 5):
    cell(ws, i, 1, i - 4, "body")
    cell(ws, i, 2, name, "body", h="left", bold=True)
    cell(ws, i, 3, code, "body")
    cell(ws, i, 4, market, "body")
    cell(ws, i, 5, industry, "body")
    cell(ws, i, 6, None, "input", fmt="0.00")
    cell(ws, i, 7, None, "input", fmt="0.00")
    cell(ws, i, 8, None, "input", fmt="0.00")
    cell(ws, i, 9, f'=IF(OR(G{i}="",H{i}=""),"",G{i}*H{i})', "formula", fmt="0.00")
    cell(ws, i, 10, f'=IF(OR(F{i}="",G{i}=""),"",G{i}/F{i})', "formula", fmt="0.00%")
    cell(ws, i, 11, None, "input", fmt="0.00")
    cell(ws, i, 12, f'=IF(OR(I{i}="",K{i}=""),"",I{i}/K{i})', "formula", fmt="0.0%")
    cell(ws, i, 13, None, "input", fmt="0.0%")
    cell(ws, i, 14, f'=IF(AND(L{i}="",M{i}=""),"",IF(IF(M{i}<>"",M{i},L{i})<=0.5,95,IF(IF(M{i}<>"",M{i},L{i})<=0.7,82,IF(IF(M{i}<>"",M{i},L{i})<=0.9,67,IF(IF(M{i}<>"",M{i},L{i})<=1,50,25)))))', "formula", fmt="0.0")
    cell(ws, i, 15, None, "input", fmt="0.00")
    cell(ws, i, 16, None, "input", fmt="0.00")
    cell(ws, i, 17, f'=IF(OR(O{i}="",P{i}=""),"",IF(P{i}<=0,"",O{i}/P{i}))', "formula", fmt="0.00")
    cell(ws, i, 18, f'=IF(Q{i}="","",IF(Q{i}<=0,95,IF(Q{i}<=1,85,IF(Q{i}<=2,72,IF(Q{i}<=3,58,IF(Q{i}<=4,42,20))))))', "formula", fmt="0.0")
    cell(ws, i, 19, None, "input", fmt="0.00")
    cell(ws, i, 20, None, "input", fmt="0.00")
    cell(ws, i, 21, f'=IF(OR(S{i}="",T{i}=""),"",S{i}/T{i})', "formula", fmt="0.00")
    cell(ws, i, 22, f'=IF(U{i}="","",IF(U{i}<=0.8,90,IF(U{i}<=1.2,78,IF(U{i}<=1.8,62,IF(U{i}<=2.5,42,20)))))', "formula", fmt="0.0")
    cell(ws, i, 23, None, "input", fmt="0.0%")
    cell(ws, i, 24, f'=IF(W{i}="","",IF(W{i}<=0.15,92,IF(W{i}<=0.3,78,IF(W{i}<=0.5,60,IF(W{i}<=0.7,40,20)))))', "formula", fmt="0.0")
    cell(ws, i, 25, None, "input", fmt="0")
    cell(ws, i, 26, f'=IF(Y{i}="","",IF(Y{i}=0,90,IF(Y{i}=1,50,20)))', "formula", fmt="0.0")
    cell(ws, i, 27, None, "input", fmt="0.0%")
    cell(ws, i, 28, None, "input", fmt="0.0")
    cell(ws, i, 29, None, "input", fmt="0.0%")
    cell(ws, i, 30, f'=IF(COUNT(N{i},R{i},V{i},X{i},Z{i},AB{i})<6,"数据不足",0.25*N{i}+0.2*R{i}+0.15*V{i}+0.15*X{i}+0.15*Z{i}+0.1*AB{i})', "formula", fmt="0.0")
    cell(ws, i, 31, None, "input")
    cell(ws, i, 32, f'=IF(AE{i}="是",40,IF(AND(Q{i}<>"",Q{i}>4),30,IF(AND(Y{i}<>"",Y{i}>=1),50,IF(AND(W{i}<>"",W{i}>0.7),50,IF(AND(AC{i}<>"",AC{i}>0.2),60,100)))))', "formula", fmt="0")
    cell(ws, i, 33, f'=IF(AE{i}="是","FCF派息率连续2年>100%，上限40",IF(AND(Q{i}<>"",Q{i}>4),"净负债/EBITDA>4，上限30",IF(AND(Y{i}<>"",Y{i}>=1),"近5年砍息，上限50",IF(AND(W{i}<>"",W{i}>0.7),"5年净利回撤>70%，上限50",IF(AND(AC{i}<>"",AC{i}>0.2),"少数股东/治理侵蚀>20%，上限60","")))))', "formula", h="left", size=9)
    cell(ws, i, 34, f'=IF(AD{i}="数据不足","",MIN(AD{i},AF{i}))', "formula", fmt="0.0")
    cell(ws, i, 35, f'=IF(AH{i}="","NA",IF(AH{i}>=80,"A",IF(AH{i}>=65,"B",IF(AH{i}>=50,"C","D"))))', "formula", bold=True)
    cell(ws, i, 36, f'=COUNTA(F{i},G{i},H{i},K{i},O{i},P{i},S{i},T{i},W{i},Y{i},AB{i},AC{i})/12', "formula", fmt="0%")
    cell(ws, i, 37, note, "note", h="left", size=9)
    ws.row_dimensions[i].height = 44

# prefill known values
# 长江电力 row 7
for c_, v, kind, fmt in [
    (6, 29.74, "input", "0.00"), (7, 0.92, "input", "0.00"), (8, 244.68217716, "input", "0.00"),
    (15, 2734.9, "input", "0.00"),
]:
    cell(ws, 7, c_, v, kind, fmt=fmt)
# 中国神华 row 8
for c_, v, kind, fmt in [
    (6, 51.99, "input", "0.00"), (7, 1.95, "input", "0.00"), (8, 216.89434304, "input", "0.00"),
    (11, 266.59, "input", "0.00"), (15, 285.9, "input", "0.00"),
]:
    cell(ws, 8, c_, v, kind, fmt=fmt)

# data validation and conditional formatting
dv_yes = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
ws.add_data_validation(dv_yes); dv_yes.add("AE5:AE11")
for val, color in [("A", LIGHT_GREEN), ("B", MID_BLUE), ("C", LIGHT_YELLOW), ("D", LIGHT_RED), ("NA", LIGHT_GRAY)]:
    ws.conditional_formatting.add("AI5:AI11", FormulaRule(formula=[f'AI5="{val}"'], fill=fill(color)))
ws.conditional_formatting.add("AG5:AG11", FormulaRule(formula=['AG5<>""'], fill=fill(LIGHT_RED), font=font(9, True, RED)))
ws.freeze_panes = "C5"
page(ws, paper=8, pages_wide=2, title_cols="A:C")

# ============================================================
# 金融红利
# ============================================================
ws = wb.create_sheet("金融红利")
ws.sheet_properties.tabColor = "2E7D6E"
widths(ws, {"A": 5, "B": 15, "C": 13, "D": 8, "E": 12, "F": 10, "G": 10, "H": 9, "I": 9, "J": 12,
            "K": 12, "L": 9, "M": 11, "N": 9, "O": 9, "P": 38, "Q": 11, "R": 44})
title(ws, 1, "金融红利 · 银行/保险引用区", 18,
      "银行不套用 FCF 六维评分；四维评分与 DDM 引用银行 skill。综合分固定权重：盈利 30 / 风险资本 30 / 估值 25 / 股息 15。保险暂标记待补。")
headers = ["序号", "标的", "代码", "市场", "类型", "盈利质量", "风险资本", "估值", "股息", "银行 skill 综合分",
           "DDM防御价值", "现价", "DDM安全边际", "股息率", "估值分", "结论", "数据日期", "备注"]
header_row(ws, 4, headers, 42)
banks = [
    ("工商银行", "601398.SH", "A股", "银行", "合理偏低；ROE 低于股权成本", "2026-10-30", "引用银行 skill 四维评分；DDM/RI 基准待填"),
    ("平安银行", "000001.SZ", "A股", "银行", "合理偏贵；ROE 低于股权成本", "2026-10-30", "DDM 8.13，综合基准 11.03；现价 11.64"),
    ("北京银行", "601169.SH", "A股", "银行", "待引用银行 skill", "", "旧 Excel 10 表，需按 11 表规范复核"),
    ("宁波银行", "002142.SZ", "A股", "银行", "待引用银行 skill", "", "旧 Excel 10 表，需按 11 表规范复核"),
]
for i, (name, code, market, typ, conclusion, data_date, note) in enumerate(banks, 5):
    cell(ws, i, 1, i - 4, "body")
    cell(ws, i, 2, name, "body", h="left", bold=True)
    cell(ws, i, 3, code, "body")
    cell(ws, i, 4, market, "body")
    cell(ws, i, 5, typ, "body")
    for c_ in range(6, 10):
        cell(ws, i, c_, None, "input", fmt="0.0")
    cell(ws, i, 10, f'=IF(COUNT(F{i}:I{i})<4,"数据不足",0.3*F{i}+0.3*G{i}+0.25*H{i}+0.15*I{i})', "formula", fmt="0.0")
    cell(ws, i, 11, None, "input", fmt="0.00")
    cell(ws, i, 12, None, "input", fmt="0.00")
    cell(ws, i, 13, f'=IF(OR(K{i}="",L{i}=""),"",1-L{i}/K{i})', "formula", fmt="0.0%")
    cell(ws, i, 14, None, "input", fmt="0.00%")
    cell(ws, i, 15, f'=IF(M{i}="","",IF(M{i}>=0.2,95,IF(M{i}>=0.05,80,IF(M{i}>=-0.05,60,IF(M{i}>=-0.2,40,20)))))', "formula", fmt="0.0")
    cell(ws, i, 16, conclusion, "body", h="left", size=9)
    cell(ws, i, 17, data_date, "body", fmt="yyyy-mm-dd")
    cell(ws, i, 18, note, "note", h="left", size=9)
    ws.row_dimensions[i].height = 40

# known dividend yields
cell(ws, 5, 14, 0.036, "input", fmt="0.00%")
cell(ws, 6, 14, 0.056, "input", fmt="0.00%")
dv_score = DataValidation(type="whole", operator="between", formula1="0", formula2="100", allow_blank=True)
ws.add_data_validation(dv_score); dv_score.add("F5:J8")
for val, color, threshold in [("A", LIGHT_GREEN, 80), ("B", MID_BLUE, 65), ("C", LIGHT_YELLOW, 50), ("D", LIGHT_RED, 0)]:
    ws.conditional_formatting.add("O5:O8", FormulaRule(formula=[f'AND(O5<>"",O5>={threshold})'], fill=fill(color)))
ws.freeze_panes = "C5"
page(ws, paper=8)

# ============================================================
# 质量防御
# ============================================================
ws = wb.create_sheet("质量防御")
ws.sheet_properties.tabColor = "6B46C1"
widths(ws, {"A": 5, "B": 15, "C": 12, "D": 8, "E": 12, "F": 9, "G": 9, "H": 11, "I": 9, "J": 10,
            "K": 10, "L": 10, "M": 12, "N": 10, "O": 10, "P": 10, "Q": 10, "R": 9, "S": 10, "T": 9,
            "U": 10, "V": 9, "W": 9, "X": 7, "Y": 34, "Z": 11, "AA": 42})
# 打印按 2 页宽拆分：标题放 A:C，保证第二页完整显示
split_title(ws, 1, "质量防御 · 五维评分", 20,
            "质量分 60% + 估值分 40%")
headers = ["序号", "标的", "代码", "市场", "类型", "ROIC", "WACC", "ROIC-WACC", "ROIC得分",
           "资本强度", "资本强度得分", "5年FCF波动", "现金流稳定性得分", "净现金/市值", "资产负债表得分",
           "回购收益率", "股东回报得分", "质量分", "DCF基准价", "现价", "安全边际", "估值分",
           "综合分", "等级", "结论", "数据日期", "备注"]
header_row(ws, 4, headers, 44)
quality = [
    ("伯克希尔", "BRK.B", "美股", "综合控股", "合理区间上沿，安全边际不足", "2026-10-30", "非股息型；现金与上市股权提供资产负债表防御"),
    ("Visa", "V", "美股", "支付", "合理偏低，安全边际不足", "2026-10-30", "PE 32.5x；股息率 0.78%；回购为主"),
    ("万事达卡", "MA", "美股", "支付", "合理偏低，安全边际不足", "2026-10-30", "PE 33.6x；股息率 0.50%；回购为主"),
]
for i, (name, code, market, typ, conclusion, data_date, note) in enumerate(quality, 5):
    cell(ws, i, 1, i - 4, "body")
    cell(ws, i, 2, name, "body", h="left", bold=True)
    cell(ws, i, 3, code, "body")
    cell(ws, i, 4, market, "body")
    cell(ws, i, 5, typ, "body")
    for c_ in (6, 7):
        cell(ws, i, c_, None, "input", fmt="0.0%")
    cell(ws, i, 8, f'=IF(OR(F{i}="",G{i}=""),"",F{i}-G{i})', "formula", fmt="0.0%")
    cell(ws, i, 9, f'=IF(H{i}="","",IF(H{i}>=0.15,95,IF(H{i}>=0.1,85,IF(H{i}>=0.05,70,IF(H{i}>=0,50,20)))))', "formula", fmt="0.0")
    cell(ws, i, 10, None, "input", fmt="0.0%")
    cell(ws, i, 11, f'=IF(J{i}="","",IF(J{i}<=0.03,95,IF(J{i}<=0.06,82,IF(J{i}<=0.1,65,IF(J{i}<=0.15,45,25)))))', "formula", fmt="0.0")
    cell(ws, i, 12, None, "input", fmt="0.0%")
    cell(ws, i, 13, f'=IF(L{i}="","",IF(L{i}<=0.1,95,IF(L{i}<=0.2,82,IF(L{i}<=0.3,65,IF(L{i}<=0.5,45,25)))))', "formula", fmt="0.0")
    cell(ws, i, 14, None, "input", fmt="0.0%")
    cell(ws, i, 15, f'=IF(N{i}="","",IF(N{i}>=0.1,95,IF(N{i}>=0.05,85,IF(N{i}>=0,70,IF(N{i}>=-0.05,55,30)))))', "formula", fmt="0.0")
    cell(ws, i, 16, None, "input", fmt="0.0%")
    cell(ws, i, 17, f'=IF(P{i}="","",IF(P{i}>=0.03,95,IF(P{i}>=0.02,85,IF(P{i}>=0.01,70,IF(P{i}>0,55,40)))))', "formula", fmt="0.0")
    cell(ws, i, 18, f'=IF(COUNT(I{i},K{i},M{i},O{i},Q{i})<5,"数据不足",0.3*I{i}+0.2*K{i}+0.2*M{i}+0.15*O{i}+0.15*Q{i})', "formula", fmt="0.0")
    cell(ws, i, 19, None, "input", fmt="0.00")
    cell(ws, i, 20, None, "input", fmt="0.00")
    cell(ws, i, 21, f'=IF(OR(S{i}="",T{i}=""),"",1-T{i}/S{i})', "formula", fmt="0.0%")
    cell(ws, i, 22, f'=IF(U{i}="","",IF(U{i}>=0.2,95,IF(U{i}>=0.05,80,IF(U{i}>=-0.05,60,IF(U{i}>=-0.2,40,20)))))', "formula", fmt="0.0")
    cell(ws, i, 23, f'=IF(AND(ISNUMBER(R{i}),ISNUMBER(V{i})),0.6*R{i}+0.4*V{i},"")', "formula", fmt="0.0")
    cell(ws, i, 24, f'=IF(ISNUMBER(W{i}),IF(W{i}>=80,"A",IF(W{i}>=65,"B",IF(W{i}>=50,"C","D"))),"")', "formula", bold=True)
    cell(ws, i, 25, conclusion, "body", h="left", size=9)
    cell(ws, i, 26, data_date, "body", fmt="yyyy-mm-dd")
    cell(ws, i, 27, note, "note", h="left", size=9)
    ws.row_dimensions[i].height = 44

# prefill DCF / price
for r, dcf, price in [(6, 424.63, 400.60), (7, 666.31, 610.31)]:
    cell(ws, r, 19, dcf, "input", fmt="0.00")
    cell(ws, r, 20, price, "input", fmt="0.00")
for val, color in [("A", LIGHT_GREEN), ("B", MID_BLUE), ("C", LIGHT_YELLOW), ("D", LIGHT_RED)]:
    ws.conditional_formatting.add("X5:X7", FormulaRule(formula=[f'X5="{val}"'], fill=fill(color)))
ws.freeze_panes = "C5"
page(ws, paper=8, pages_wide=2, title_cols="A:C")

# ============================================================
# 总回报
# ============================================================
ws = wb.create_sheet("总回报")
ws.sheet_properties.tabColor = "2F5597"
widths(ws, {"A": 5, "B": 15, "C": 13, "D": 8, "E": 8, "F": 9, "G": 10, "H": 9, "I": 9, "J": 9,
            "K": 10, "L": 10, "M": 9, "N": 8, "O": 9, "P": 9, "Q": 10, "R": 9, "S": 9, "T": 8,
            "U": 11, "V": 46})
title(ws, 1, "总回报 · 税后股息 + 价格 + 汇率 − 10Y", 22,
      "税率、价格回报、汇率必须填写后，总回报才计算；空缺不是 0。10Y 为本地 10 年期国债收益率。")
headers = ["序号", "标的", "代码", "市场", "币种", "现价", "2026E DPS", "股息率", "本地10Y", "股息-10Y",
           "预期DPS增长", "预期价格回报", "汇率变动", "税率", "税后股息率", "税后总回报", "总回报-10Y",
           "估值基准价", "安全边际", "估值分", "数据日期", "备注"]
header_row(ws, 4, headers, 44)
all_names = [
    ("工商银行", "601398.SH", "A股", "CNY", 0.036, 0.0162, None, None, "股息率 3.6%；价格/DPS 待填"),
    ("平安银行", "000001.SZ", "A股", "CNY", None, 0.0162, 11.64, 0.65, "综合基准 11.03；税率待填"),
    ("美的集团", "000333.SZ", "A股", "CNY", 0.049, 0.0162, None, None, "股息率约4.9%；价格/DPS待填"),
    ("贵州茅台", "600519.SH", "A股", "CNY", 0.038, 0.0162, None, None, "股息率约3.8%；价格/DPS待填"),
    ("长江电力", "600900.SH", "A股", "CNY", None, 0.0162, 29.74, 0.92, "DCF/DDM 基准 27.59；税率待填"),
    ("中国神华", "601088.SH", "A股", "CNY", None, 0.0162, 51.99, 1.95, "三法基准 36.64；税率待填"),
    ("可口可乐", "KO", "美股", "USD", None, 0.0490, None, None, "PE 27x；价格/DPS待填"),
    ("中国移动", "600941.SH", "A股", "CNY", None, 0.0162, None, None, "待研究"),
    ("宁沪高速", "600377.SH", "A股", "CNY", None, 0.0162, None, None, "待研究"),
    ("伯克希尔", "BRK.B", "美股", "USD", None, 0.0490, None, None, "非股息型；价格/DCF待填"),
    ("Visa", "V", "美股", "USD", 0.0078, 0.0490, 400.60, None, "股息率0.78%；DCF 424.63"),
    ("万事达卡", "MA", "美股", "USD", 0.0050, 0.0490, 610.31, None, "股息率0.50%；DCF 666.31"),
]
# baseline values for certain names
baseline = {6: 11.03, 9: 27.59, 10: 36.64, 15: 424.63, 16: 666.31}
for i, (name, code, market, ccy, yield_, ten_y, price, dps, note) in enumerate(all_names, 5):
    cell(ws, i, 1, i - 4, "body")
    cell(ws, i, 2, name, "body", h="left", bold=True)
    cell(ws, i, 3, code, "body")
    cell(ws, i, 4, market, "body")
    cell(ws, i, 5, ccy, "body")
    cell(ws, i, 6, price, "input", fmt="0.00")
    cell(ws, i, 7, dps, "input", fmt="0.00")
    if yield_ is None:
        # 价格/DPS 都有时自动计算；两者缺一时留空，不用 0 代替
        h_formula = f'=IF(AND(F{i}<>"",G{i}<>""),G{i}/F{i},"")'
    else:
        # 只有已知股息率、缺价格/DPS 时，先显示来源股息率；补齐后自动改为 DPS/价格
        h_formula = f'=IF(AND(F{i}<>"",G{i}<>""),G{i}/F{i},{yield_})'
    cell(ws, i, 8, h_formula, "formula", fmt="0.00%")
    cell(ws, i, 9, ten_y, "input", fmt="0.00%")
    cell(ws, i, 10, f'=IF(OR(H{i}="",I{i}=""),"",H{i}-I{i})', "formula", fmt="0.00%")
    cell(ws, i, 11, None, "input", fmt="0.0%")
    cell(ws, i, 12, None, "input", fmt="0.0%")
    cell(ws, i, 13, None, "input", fmt="0.0%")
    cell(ws, i, 14, None, "input", fmt="0.0%")
    cell(ws, i, 15, f'=IF(OR(H{i}="",N{i}=""),"",H{i}*(1-N{i}))', "formula", fmt="0.00%")
    cell(ws, i, 16, f'=IF(OR(O{i}="",L{i}="",M{i}=""),"",(1+O{i})*(1+L{i})*(1+M{i})-1)', "formula", fmt="0.00%")
    cell(ws, i, 17, f'=IF(OR(P{i}="",I{i}=""),"",P{i}-I{i})', "formula", fmt="0.00%")
    cell(ws, i, 18, baseline.get(i), "input", fmt="0.00")
    cell(ws, i, 19, f'=IF(OR(R{i}="",F{i}=""),"",1-F{i}/R{i})', "formula", fmt="0.0%")
    cell(ws, i, 20, f'=IF(AND(Q{i}="",S{i}=""),"",IF(AND(Q{i}<>"",S{i}<>""),0.5*IF(Q{i}>=0.06,95,IF(Q{i}>=0.04,82,IF(Q{i}>=0.02,68,IF(Q{i}>=0,52,IF(Q{i}>=-0.02,38,20)))))+0.5*IF(S{i}>=0.2,95,IF(S{i}>=0.05,80,IF(S{i}>=-0.05,60,IF(S{i}>=-0.2,40,20)))),IF(S{i}<>"",IF(S{i}>=0.2,95,IF(S{i}>=0.05,80,IF(S{i}>=-0.05,60,IF(S{i}>=-0.2,40,20)))),IF(Q{i}>=0.06,95,IF(Q{i}>=0.04,82,IF(Q{i}>=0.02,68,IF(Q{i}>=0,52,IF(Q{i}>=-0.02,38,20))))))))', "formula", fmt="0.0")
    cell(ws, i, 21, "2026-10-30", "body", fmt="yyyy-mm-dd")
    cell(ws, i, 22, note, "note", h="left", size=9)
    ws.row_dimensions[i].height = 36

for val, color, threshold in [("A", LIGHT_GREEN, 80), ("B", MID_BLUE, 65), ("C", LIGHT_YELLOW, 50), ("D", LIGHT_RED, 0)]:
    ws.conditional_formatting.add("T5:T16", FormulaRule(formula=[f'AND(T5<>"",T5>={threshold})'], fill=fill(color)))
ws.conditional_formatting.add("Q5:Q16", CellIsRule(operator="lessThan", formula=["0"], font=font(10, True, RED)))
ws.conditional_formatting.add("Q5:Q16", CellIsRule(operator="greaterThan", formula=["0"], font=font(10, True, GREEN)))
ws.freeze_panes = "C5"
page(ws, paper=8)

# ============================================================
# 组合监控
# ============================================================
ws = wb.create_sheet("组合监控")
ws.sheet_properties.tabColor = "6B46C1"
widths(ws, {"A": 16, "B": 14, "C": 14, "D": 10, "E": 10, "F": 9, "G": 9, "H": 11, "I": 10, "J": 9,
            "K": 10, "L": 16, "M": 34})
title(ws, 1, "组合监控 · 从单票到配置", 13,
      "示例约束，不是投资建议。权重、行业、利率敏感度、相关性、触发条件一起看。")
headers = ["标的", "类型", "行业", "目标权重", "当前权重", "股息率", "总回报", "利率敏感度", "相关性桶", "流动性", "最大回撤", "风险标签", "备注"]
header_row(ws, 4, headers, 40)
for i, (name, typ, industry) in enumerate([
    ("工商银行", "金融红利", "银行"), ("平安银行", "金融红利", "银行"),
    ("美的集团", "一般红利", "家电"), ("贵州茅台", "一般红利", "消费"),
    ("长江电力", "一般红利/类债", "公用事业"), ("中国神华", "一般红利", "煤炭"),
    ("可口可乐", "一般红利", "消费"), ("中国移动", "一般红利", "运营商"),
    ("宁沪高速", "一般红利", "高速公路"), ("伯克希尔", "质量防御", "控股"),
    ("Visa", "质量防御", "支付"), ("万事达卡", "质量防御", "支付"),
], 5):
    cell(ws, i, 1, name, "body", h="left", bold=True)
    cell(ws, i, 2, typ, "body")
    cell(ws, i, 3, industry, "body")
    cell(ws, i, 4, None, "input", fmt="0.0%")
    cell(ws, i, 5, None, "input", fmt="0.0%")
    cell(ws, i, 6, f"='总回报'!H{i}", "formula", fmt="0.00%")
    cell(ws, i, 7, f"='总回报'!P{i}", "formula", fmt="0.00%")
    cell(ws, i, 8, None, "input")
    cell(ws, i, 9, None, "input")
    cell(ws, i, 10, None, "input")
    cell(ws, i, 11, None, "input", fmt="0.0%")
    cell(ws, i, 12, None, "input")
    cell(ws, i, 13, None, "note", h="left", size=9)
    ws.row_dimensions[i].height = 24

# summary
cell(ws, 18, 1, "汇总", "section", h="left")
ws.merge_cells("A18:M18")
summary = [
    ("红利/类债合计", '=SUMIF($B$5:$B$16,"*红利*",$D$5:$D$16)', "0.0%"),
    ("质量防御合计", '=SUMIF($B$5:$B$16,"质量防御",$D$5:$D$16)', "0.0%"),
    ("最大单一权重", "=MAX(D5:D16)", "0.0%"),
    ("加权股息率", '=IF(SUM(D5:D16)=0,"",SUMPRODUCT(D5:D16,F5:F16)/SUM(D5:D16))', "0.00%"),
    ("加权总回报", '=IF(SUM(D5:D16)=0,"",SUMPRODUCT(D5:D16,G5:G16)/SUM(D5:D16))', "0.00%"),
    ("高利率敏感度占比", '=IF(SUM(D5:D16)=0,"",SUMIF(H5:H16,"高",D5:D16)/SUM(D5:D16))', "0.0%"),
    ("权重合计检查", '=IF(ABS(SUM(D5:D16)-1)<0.001,"100%","未到 100%")', "General"),
]
for r, (label, formula, fmt) in enumerate(summary, 19):
    cell(ws, r, 1, label, "body", h="left", bold=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    cell(ws, r, 3, formula, "formula", fmt=fmt, h="left")

# sector concentration
cell(ws, 28, 1, "行业集中度", "section", h="left")
ws.merge_cells("A28:M28")
header_row(ws, 29, ["行业", "目标权重"])
for r, sector in enumerate(["银行", "煤炭", "公用事业", "消费", "支付", "控股", "家电", "运营商", "高速公路"], 30):
    cell(ws, r, 1, sector, "body", h="left")
    cell(ws, r, 2, f'=SUMIF($C$5:$C$16,A{r},$D$5:$D$16)', "formula", fmt="0.0%", h="left")

# triggers
cell(ws, 41, 1, "触发条件", "section", h="left")
ws.merge_cells("A41:M41")
header_row(ws, 42, ["检查项", "状态", "动作"])
for c_ in (4, 5, 6):
    cell(ws, 42, c_, None, "header")
ws.merge_cells("C42:F42")
triggers = [
    ("砍息 / 派息下调", "正常", "回到单票报告复核"),
    ("FCF 派息率连续 2 年 >100%", "待核", "检查 FCF 口径与资本开支"),
    ("净负债/EBITDA >4x", "待核", "检查资产负债表与评级"),
    ("现价高于基准价值 20%", "待核", "减仓/不追高讨论"),
    ("核心逻辑破坏", "正常", "重新写论文"),
    ("数据超过 90 天", "待核", "更新价格/财报/分红"),
]
for r, (item, status, action) in enumerate(triggers, 43):
    cell(ws, r, 1, item, "body", h="left")
    cell(ws, r, 2, status, "input")
    cell(ws, r, 3, action, "note", h="left", size=9)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)

dv_sens = DataValidation(type="list", formula1='"高,中,低"', allow_blank=True)
ws.add_data_validation(dv_sens); dv_sens.add("H5:J16")
dv_status = DataValidation(type="list", formula1='"正常,触发,待核"', allow_blank=True)
ws.add_data_validation(dv_status); dv_status.add("B43:B48")
ws.freeze_panes = "A5"
# 主表+汇总一页；行业集中度+触发条件另起一页
ws.row_breaks.append(Break(id=27))
page(ws, paper=9)

# ============================================================
# 数据与来源
# ============================================================
ws = wb.create_sheet("数据与来源")
ws.sheet_properties.tabColor = "5B6B7C"
widths(ws, {"A": 15, "B": 28, "C": 14, "D": 10, "E": 24, "F": 11, "G": 11, "H": 30, "I": 10, "J": 34})
title(ws, 1, "数据与来源 · 证据账本", 10,
      "每个关键数字至少记录：来源、发布日、期间、原始值/单位、状态。状态可选：硬数据/软数据/假设/计算/待核。")
headers = ["标的", "字段", "数值", "单位", "来源", "发布日", "期间", "链接/附件", "状态", "备注"]
header_row(ws, 4, headers, 38)
sources = [
    ("长江电力", "当前价", 29.74, "元/股", "腾讯行情", "2026-10-30", "2026-10-30", "", "硬数据", "收盘价口径"),
    ("长江电力", "2026E DPS", 0.92, "元/股", "公司公告/模型", "2026-10-30", "2026E", "", "假设", "待按最新公告复核"),
    ("长江电力", "净负债", 2734.9, "亿元", "2026 中报", "2026-08-31", "2026H1", "", "硬数据", "有息负债−货币资金口径待核"),
    ("长江电力", "FY2025 FCF", None, "亿元", "FY2025 年报", "", "FY2025", "", "待核", "经营现金流−购建固定资产等支付的现金"),
    ("中国神华", "FY2025 FCF", 266.59, "亿元", "FY2025 年报 / 煤炭 skill", "2026-10-30", "FY2025", "", "硬数据", "煤炭 skill 锁死口径"),
    ("中国神华", "2026E DPS", 1.95, "元/股", "模型", "2026-10-30", "2026E", "", "假设", "需按最新公告复核"),
    ("中国神华", "净负债", 285.9, "亿元", "2026 中报", "2026-08-31", "2026H1", "", "硬数据", "有息负债−货币资金口径待核"),
    ("平安银行", "当前价", 11.64, "元/股", "腾讯行情", "2026-10-30", "2026-10-30", "", "硬数据", "收盘价口径"),
    ("平安银行", "综合基准", 11.03, "元/股", "巴菲特框架研究", "2026-10-30", "2026-10-30", "", "计算", "RI 70% + DDM 30%"),
    ("工商银行", "股息率", 0.036, "%", "巴菲特框架研究", "2026-10-29", "2026E", "", "计算", "前瞻股息率"),
    ("Visa", "当前价", 400.60, "美元/股", "腾讯行情", "2026-10-29", "2026-10-29", "", "硬数据", "收盘价口径"),
    ("Visa", "DCF 基准", 424.63, "美元/股", "巴菲特框架研究", "2026-10-29", "2026-10-29", "", "计算", "DCF 基准情景"),
    ("万事达卡", "当前价", 610.31, "美元/股", "腾讯行情", "2026-10-29", "2026-10-29", "", "硬数据", "收盘价口径"),
    ("万事达卡", "DCF 基准", 666.31, "美元/股", "巴菲特框架研究", "2026-10-29", "2026-10-29", "", "计算", "DCF 基准情景"),
    ("中国 10Y", "10 年期国债收益率", 0.0162, "%", "WebSearch", "2026-10-30", "2026-10-30", "", "硬数据", "本地无风险利率"),
    ("美国 10Y", "10 年期国债收益率", 0.0490, "%", "WebSearch", "2026-10-30", "2026-10-30", "", "硬数据", "本地无风险利率"),
]
for r, row in enumerate(sources, 5):
    for c_, v in enumerate(row, 1):
        kind = "input" if c_ in (3, 6, 7, 8, 9) else "body"
        fmt = "0.00%" if c_ == 3 and isinstance(v, float) and v < 1 else ("0.00" if isinstance(v, float) else None)
        cell(ws, r, c_, v, kind, fmt=fmt, h="left" if c_ in (1, 2, 5, 8, 10) else "center", size=9)
    ws.row_dimensions[r].height = 30

dv_status2 = DataValidation(type="list", formula1='"硬数据,软数据,假设,计算,待核"', allow_blank=True)
ws.add_data_validation(dv_status2); dv_status2.add("I5:I40")
ws.freeze_panes = "A5"
page(ws)

# ============================================================
# 检查
# ============================================================
ws = wb.create_sheet("检查")
ws.sheet_properties.tabColor = "B42318"
widths(ws, {"A": 38, "B": 22, "C": 12, "D": 52})
title(ws, 1, "自动检查", 4, "公式结果需在 Excel 中重算后查看；本页只做结构检查，不替代人工复核。")
header_row(ws, 4, ["检查项", "结果", "状态", "说明"])
checks = [
    ("六维权重合计", "=SUM(0.25,0.2,0.15,0.15,0.15,0.1)", '=IF(ABS(B5-1)<0.0001,"OK","CHECK")', "应为 100%"),
    ("质量分越界数量", "=COUNTIF('分红安全'!AD5:AD11,\">100\")+COUNTIF('分红安全'!AD5:AD11,\"<0\")", '=IF(B6=0,"OK","CHECK")', "质量分应在 0–100"),
    ("最终质量分越界", "=COUNTIF('分红安全'!AH5:AH11,\">100\")+COUNTIF('分红安全'!AH5:AH11,\"<0\")", '=IF(B7=0,"OK","CHECK")', "硬约束后的分数应在 0–100"),
    ("硬约束上限越界", "=COUNTIF('分红安全'!AF5:AF11,\">100\")+COUNTIF('分红安全'!AF5:AF11,\"<0\")", '=IF(B8=0,"OK","CHECK")', "上限应为 0–100"),
    ("数据完整度 <80%（含待研究）", "=COUNTIF('分红安全'!AJ5:AJ11,\"<0.8\")", '=IF(B9=0,"OK","CHECK")', "模板阶段为正常；完成研究后应清零"),
    ("重复代码", "=SUMPRODUCT(--(COUNTIF('评分总表'!C5:C16,'评分总表'!C5:C16)>1))", '=IF(B10=0,"OK","CHECK")', "同一代码不应重复"),
    ("缺少数据日期", "=COUNTBLANK('评分总表'!P5:P16)", '=IF(B11=0,"OK","CHECK")', "待研究行可留空"),
    ("总回报估值分越界", "=COUNTIF('总回报'!T5:T16,\">100\")+COUNTIF('总回报'!T5:T16,\"<0\")", '=IF(B12=0,"OK","CHECK")', "估值分应在 0–100"),
    ("组合权重合计", "=SUM('组合监控'!D5:D16)", '=IF(ABS(B13-1)<0.001,"OK","CHECK")', "建仓前为 0 属正常；建仓后应为 100%"),
    ("组合权重为负", "=COUNTIF('组合监控'!D5:D16,\"<0\")", '=IF(B14=0,"OK","CHECK")', "不允许负权重"),
    ("免责声明", "本表只做研究分级，不构成投资建议", "OK", "输出前保留"),
]
for r, (item, formula, status, note) in enumerate(checks, 5):
    cell(ws, r, 1, item, "body", h="left", bold=True)
    cell(ws, r, 2, formula, "formula", h="center")
    cell(ws, r, 3, status, "formula", bold=True)
    cell(ws, r, 4, note, "note", h="left", size=9)
    ws.row_dimensions[r].height = 28

ws.conditional_formatting.add("C5:C15", FormulaRule(formula=['C5="OK"'], fill=fill(LIGHT_GREEN)))
ws.conditional_formatting.add("C5:C15", FormulaRule(formula=['C5="CHECK"'], fill=fill(LIGHT_RED)))
page(ws, paper=9, pages_high=1)

# page setup / print titles
for s in wb.worksheets:
    s.sheet_view.showGridLines = False
    s.print_title_rows = "1:4"
# 说明页第二页不需要重复“定位”段；组合监控第二页不重复主表表头
wb["说明"].print_title_rows = "1:2"
wb["组合监控"].print_title_rows = "1:3"
# 横向拆成 2 页的明细表：第二页只重复表头，不重复会被截断的标题
wb["分红安全"].print_title_rows = "4:4"
wb["质量防御"].print_title_rows = "4:4"

# 检查页数字格式
wb["检查"]["B5"].number_format = "0%"
wb["检查"]["B13"].number_format = "0.0%"

wb.save(OUT)
print("saved")
