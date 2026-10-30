# -*- coding: utf-8 -*-
"""Build Alphabet (GOOGL) Buffett-framework valuation workbook.

All figures in USD billions unless noted. Data cut-off: 2026-10-29.
Official financials: FY2025 10-K (ended 2025-12-31) and Q3 2026 earnings
release (quarter ended 2026-09-30, 8-K filed 2026-10-28). Q3 2026 10-Q not
yet filed. Reference price: $291.50 after-hours on 2026-10-28; last regular
close $324.89.
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
XLSX = BASE / "GOOGL-估值模型-2026-10-29.xlsx"

# ----------------------------------------------------------------------------
# Inputs and computed outputs
# ----------------------------------------------------------------------------
PRICE_AH = 291.50          # 2026-10-28 after-hours indicative
PRICE_CLOSE = 324.89       # 2026-10-28 regular close
SHARES = 12.294            # billion, 2026-09-30
CASH = 265.098             # cash + marketable securities
DEBT = 125.102             # long-term debt
PREF = 19.000              # mandatory convertible preferred (liquidation)
GROSS_NET_CASH = CASH - DEBT
NET_CASH_COMMON = GROSS_NET_CASH - PREF
MCAP_AH = PRICE_AH * SHARES
MCAP_CLOSE = PRICE_CLOSE * SHARES
EV_AH = MCAP_AH - NET_CASH_COMMON
EV_CLOSE = MCAP_CLOSE - NET_CASH_COMMON

TTM_REV = 471.098
TTM_EBIT = 156.976
TTM_NOPAT = 127.15
TTM_FCF = 19.747
TTM_EPS_ADJ = 10.34
REV_2026E = 499.0
REV_2027E = 620.0
EPS_2026E = 11.89
EPS_2027E = 15.19
TAX = 0.19
NWC = 0.02

SCEN = {
    "bear": dict(
        label="悲观", wacc=0.093, gt=0.030,
        growth=[0.21, 0.19, 0.13, 0.10, 0.08, 0.06, 0.05, 0.04, 0.03, 0.03],
        ebit=[0.305] * 10,
        da=[0.060, 0.070, 0.080, 0.085, 0.085, 0.085, 0.085, 0.085, 0.085, 0.085],
        capex=[0.410, 0.500, 0.420, 0.340, 0.280, 0.240, 0.220, 0.200, 0.180, 0.160],
    ),
    "base": dict(
        label="基准", wacc=0.088, gt=0.035,
        growth=[0.24, 0.23, 0.17, 0.14, 0.11, 0.09, 0.08, 0.07, 0.06, 0.05],
        ebit=[0.320, 0.330, 0.335, 0.335, 0.340, 0.340, 0.340, 0.340, 0.340, 0.340],
        da=[0.060, 0.070, 0.080, 0.085, 0.085, 0.085, 0.085, 0.085, 0.085, 0.085],
        capex=[0.400, 0.490, 0.400, 0.300, 0.240, 0.200, 0.180, 0.170, 0.160, 0.150],
    ),
    "bull": dict(
        label="乐观", wacc=0.080, gt=0.040,
        growth=[0.26, 0.27, 0.20, 0.17, 0.14, 0.12, 0.10, 0.09, 0.08, 0.07],
        ebit=[0.325, 0.340, 0.350, 0.355, 0.360, 0.360, 0.360, 0.360, 0.360, 0.360],
        da=[0.060, 0.070, 0.080, 0.085, 0.085, 0.085, 0.085, 0.085, 0.085, 0.085],
        capex=[0.390, 0.470, 0.370, 0.270, 0.210, 0.175, 0.155, 0.140, 0.130, 0.120],
    ),
}


def dcf(s):
    rev = TTM_REV
    prev = rev
    pv = 0.0
    rows = []
    for t in range(10):
        rev = prev * (1 + s["growth"][t])
        ebit = rev * s["ebit"][t]
        nopat = ebit * (1 - TAX)
        da = rev * s["da"][t]
        capex = rev * s["capex"][t]
        dnwc = (rev - prev) * NWC
        fcf = nopat + da - capex - dnwc
        df = 1 / (1 + s["wacc"]) ** (t + 1)
        rows.append(dict(year=t + 1, growth=s["growth"][t], rev=rev, ebit=ebit,
                         nopat=nopat, da=da, capex=capex, dnwc=dnwc,
                         fcf=fcf, df=df, pv=fcf * df))
        pv += fcf * df
        prev = rev
    tv = rows[-1]["fcf"] * (1 + s["gt"]) / (s["wacc"] - s["gt"])
    pvtv = tv / (1 + s["wacc"]) ** 10
    ev = pv + pvtv
    equity = ev + NET_CASH_COMMON
    return dict(rows=rows, pv=pv, tv=tv, pvtv=pvtv, ev=ev,
                equity=equity, ps=equity / SHARES,
                tv_share=pvtv / ev, term_fcf=rows[-1]["fcf"],
                term_rev=rows[-1]["rev"])


RESULTS = {k: dcf(v) for k, v in SCEN.items()}


def bisect(f, lo, hi, tol=1e-10):
    flo = f(lo)
    for _ in range(300):
        mid = (lo + hi) / 2
        fm = f(mid)
        if abs(fm) < tol:
            return mid
        if flo * fm < 0:
            hi = mid
        else:
            lo = mid
            flo = fm
    return (lo + hi) / 2


def r0(x):
    """Round half-up for display tables (avoid Python banker's rounding)."""
    return int(x + 0.5 + 1e-6) if x >= 0 else -int(-x + 0.5 + 1e-6)


def irr_at(price, s, rows, term_fcf):
    def npv(r):
        p = sum(row["fcf"] / (1 + r) ** row["year"] for row in rows)
        tv = term_fcf * (1 + s["gt"]) / (r - s["gt"])
        return p + tv / (1 + r) ** 10 + NET_CASH_COMMON - price * SHARES
    return bisect(npv, s["gt"] + 1e-4, 0.50)


IRR_AH = {k: irr_at(PRICE_AH, v, RESULTS[k]["rows"], RESULTS[k]["term_fcf"]) for k, v in SCEN.items()}
IRR_CLOSE = {k: irr_at(PRICE_CLOSE, v, RESULTS[k]["rows"], RESULTS[k]["term_fcf"]) for k, v in SCEN.items()}

# WACC sensitivity on base case
WACC_GRID = [0.080, 0.083, 0.086, 0.088, 0.090, 0.093, 0.096, 0.100]
WACC_PS = []
for w in WACC_GRID:
    s = dict(SCEN["base"])
    s["wacc"] = w
    WACC_PS.append(dcf(s)["ps"])

# Terminal capex sensitivity on base case
CAPEX_GRID = [0.130, 0.140, 0.150, 0.160, 0.170]
CAPEX_PS = []
for c in CAPEX_GRID:
    s = dict(SCEN["base"])
    s["capex"] = list(SCEN["base"]["capex"][:-1]) + [c]
    CAPEX_PS.append(dcf(s)["ps"])

# Terminal EBIT margin sensitivity on base case
MARGIN_GRID = [0.320, 0.330, 0.340, 0.350, 0.360]
MARGIN_PS = []
for m in MARGIN_GRID:
    s = dict(SCEN["base"])
    s["ebit"] = list(SCEN["base"]["ebit"][:-1]) + [m]
    MARGIN_PS.append(dcf(s)["ps"])


def value_const_g(g, base=SCEN["base"]):
    rev = TTM_REV
    prev = rev
    pv = 0.0
    last = 0.0
    for t in range(10):
        rev = prev * (1 + g)
        nopat = rev * base["ebit"][t] * (1 - TAX)
        fcf = nopat + rev * base["da"][t] - rev * base["capex"][t] - (rev - prev) * NWC
        pv += fcf / (1 + base["wacc"]) ** (t + 1)
        last = fcf
        prev = rev
    tv = last * (1 + base["gt"]) / (base["wacc"] - base["gt"])
    return pv + tv / (1 + base["wacc"]) ** 10 + NET_CASH_COMMON


IMPLIED_G_AH = bisect(lambda g: value_const_g(g) - PRICE_AH * SHARES, -0.2, 0.4)
IMPLIED_G_CLOSE = bisect(lambda g: value_const_g(g) - PRICE_CLOSE * SHARES, -0.2, 0.4)

# SOTP cross-check (TTM segment EBIT; net cash + 50% of non-marketable securities)
SERVICES_EBIT = 159.422
CLOUD_EBIT = 31.846
OTHER_BETS = -9.563
CORP_COST = -24.729
INVESTMENTS_HAIRCUT = 135.230 * 0.50
SOTP = []
for sm in [17, 18, 19, 20]:
    for cm in [25, 30, 35]:
        eq = (SERVICES_EBIT * sm + CLOUD_EBIT * cm + OTHER_BETS
              + CORP_COST * 15 + NET_CASH_COMMON + INVESTMENTS_HAIRCUT)
        SOTP.append((sm, cm, eq, eq / SHARES))

# ----------------------------------------------------------------------------
# Styling helpers
# ----------------------------------------------------------------------------
FONT = "Microsoft YaHei"
C_DARK = "1F2933"
C_LIGHT = "F7F4EC"
C_CREAM = "E8E2D5"
C_RED = "B23A2E"
C_GREEN = "2F6B4F"
C_GOLD = "C9A227"
C_BLUE = "1A73E8"
C_GREY = "5B6770"
C_YELLOW = "FFF2CC"
C_WHITE = "FFFFFF"

thin = Side(style="thin", color="D8D2C4")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def title(ws, cell, text, size=16):
    ws[cell] = text
    ws[cell].font = Font(name=FONT, size=size, bold=True, color=C_DARK)


def subtitle(ws, cell, text):
    ws[cell] = text
    ws[cell].font = Font(name=FONT, size=10, color=C_GREY)


def section(ws, row, text, ncols=8):
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=FONT, size=11, bold=True, color=C_DARK)
    for c in range(1, ncols + 1):
        ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor=C_CREAM)
        ws.cell(row=row, column=c).border = BORDER


def header_row(ws, row, values):
    for j, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=j, value=v)
        cell.font = Font(name=FONT, size=9, bold=True, color=C_WHITE)
        cell.fill = PatternFill("solid", fgColor=C_DARK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def body_cell(ws, row, col, value, fmt=None, bold=False, align="right", color=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(name=FONT, size=9, bold=bold, color=color or C_DARK)
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.border = BORDER
    if fmt:
        cell.number_format = fmt
    return cell


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


# ----------------------------------------------------------------------------
# Workbook
# ----------------------------------------------------------------------------
wb = Workbook()

# --- Sheet 1: 结论 -----------------------------------------------------------
ws = wb.active
ws.title = "结论"
ws.sheet_view.showGridLines = False
widths(ws, {"A": 28, "B": 16, "C": 18, "D": 20, "E": 14})

title(ws, "A1", "Alphabet（GOOGL）巴菲特框架估值")
subtitle(ws, "A2", "研究日 2026-10-29｜盘后参考价 $291.50｜10/28 收盘 $324.89｜最新财报：2026 Q3（2026-09-30）")
subtitle(ws, "A3", "Q3 2026 10-Q 尚未提交；财务数据来自 8-K 业绩稿（未经审计）。市值按 12.294B 股计算。")

section(ws, 5, "一句话结论", 5)
ws["A6"] = "好生意，但 AI 资本开支把自由现金流打到负值；价格大致合理，安全边际不足。结论：观察，等 FCF 拐点。"
ws["A6"].font = Font(name=FONT, size=11, bold=True, color=C_RED)
ws["A6"].alignment = Alignment(wrap_text=True, vertical="center")
ws.merge_cells("A6:E6")
ws.row_dimensions[6].height = 38

section(ws, 8, "关键数据（十亿美元 / 倍 / 百分比）", 5)
header_row(ws, 9, ["项目", "数值", "备注"])
rows = [
    ("盘后参考价（美元）", PRICE_AH, "2026-10-28 盘后，约 -10%", "$#,##0.00"),
    ("10/28 收盘价（美元）", PRICE_CLOSE, "财报前收盘", "$#,##0.00"),
    ("期末股数（十亿股）", SHARES, "2026-09-30，A+B+C", "0.000"),
    ("市值 @ 盘后（十亿美元）", "=B10*B12", "价格 × 股数", "#,##0.0"),
    ("净现金（十亿美元）", round(NET_CASH_COMMON, 3), "现金及证券 − 债务 − 优先股", "#,##0.0"),
    ("企业价值 @ 盘后", "=B13-B14", "市值 − 净现金", "#,##0.0"),
    ("TTM 收入", TTM_REV, "2025Q4–2026Q3", "#,##0.0"),
    ("TTM 经营利润", TTM_EBIT, "经营利润，剔除 OI&E", "#,##0.0"),
    ("TTM NOPAT", TTM_NOPAT, "经营利润 × (1−19%)", "#,##0.0"),
    ("TTM 自由现金流", TTM_FCF, "CFO − capex", "#,##0.0"),
    ("2026E 收入", REV_2026E, "一致预期约 498–500", "#,##0.0"),
    ("2027E 收入", REV_2027E, "一致预期约 612–628", "#,##0.0"),
    ("2026E 调整后 EPS（美元）", EPS_2026E, "剔除投资损益", "0.00"),
    ("2027E 调整后 EPS（美元）", EPS_2027E, "剔除一次性收益", "0.00"),
    ("EV / TTM EBIT", "=B15/B17", "盘后价", '0.0"x"'),
    ("P/E 2026E", "=B10/B22", "盘后价 / 调整后 EPS", '0.0"x"'),
    ("P/E 2027E", "=B10/B23", "盘后价 / 调整后 EPS", '0.0"x"'),
    ("FCF 收益率", "=B19/B13", "TTM FCF / 市值", "0.00%"),
    ("股息率", "=0.88/B10", "年化股息 $0.88", "0.00%"),
]
for i, (label, value, note, fmt) in enumerate(rows, start=10):
    body_cell(ws, i, 1, label, align="left")
    body_cell(ws, i, 2, value, fmt=fmt, bold=isinstance(value, str) and value.startswith("="))
    body_cell(ws, i, 3, note, align="left")

section(ws, 31, "三情景 DCF（FCFF，WACC 折现，加净现金）", 5)
header_row(ws, 32, ["情景", "每股价值（美元）", "较盘后价", "安全边际（1−价/值）", "10 年 IRR"])
for i, k in enumerate(["bear", "base", "bull"], start=33):
    r = RESULTS[k]
    body_cell(ws, i, 1, SCEN[k]["label"], align="left", bold=(k == "base"))
    body_cell(ws, i, 2, round(r["ps"], 2), fmt="$#,##0.00", bold=(k == "base"))
    body_cell(ws, i, 3, f"=B{i}/$B$10-1", fmt="0.0%")
    body_cell(ws, i, 4, f"=1-$B$10/B{i}", fmt="0.0%")
    body_cell(ws, i, 5, round(IRR_AH[k], 4), fmt="0.00%")
body_cell(ws, 36, 1, "盘后价", align="left", bold=True, color=C_RED)
body_cell(ws, 36, 2, PRICE_AH, fmt="$#,##0.00", bold=True, color=C_RED)
body_cell(ws, 36, 3, 0, fmt="0.0%")
body_cell(ws, 36, 4, 0, fmt="0.0%")
body_cell(ws, 36, 5, "—", align="center")

section(ws, 38, "结论与观察条件", 5)
notes = [
    "基准价值约 270 美元，较盘后价低 7%；乐观 566 美元，悲观 126 美元。",
    "当前 EV 约 3.46 万亿美元；按基准假设，隐含未来 10 年收入 CAGR 约 13.4%（基准 12.2%）。",
    "FCF 是核心变量：2026E capex 195–205B，2027E 显著增加（媒体口径 >300B）。",
    "若要求 20% 安全边际：基准观察价约 216 美元；乐观情景约 453 美元。",
    "盘后价 291.50 高于基准价值；10/28 收盘 324.89 更高。结论是“观察”，不是“卖出”。",
    "这不是投资建议；估值对 capex 正常化速度、WACC 和 Cloud 利润率高度敏感。",
]
for i, text in enumerate(notes, start=39):
    ws.cell(row=i, column=1, value="• " + text).font = Font(name=FONT, size=10, color=C_DARK)
    ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=5)
    ws.row_dimensions[i].height = 22

chart = BarChart()
chart.type = "bar"
chart.style = 10
chart.title = "每股价值：三情景 vs 盘后价（美元）"
chart.y_axis.title = "美元/股"
data = Reference(ws, min_col=2, min_row=32, max_row=36)
cats = Reference(ws, min_col=1, min_row=33, max_row=36)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.height = 7.5
chart.width = 13
chart.legend = None
ws.add_chart(chart, "G9")

# --- Sheet 2: 历史财务 -------------------------------------------------------
ws = wb.create_sheet("历史财务")
ws.sheet_view.showGridLines = False
widths(ws, {"A": 28, "B": 11, "C": 11, "D": 11, "E": 11, "F": 11, "G": 12, "H": 30})
title(ws, "A1", "Alphabet 五年财务与现金流（十亿美元，除每股数据）")
subtitle(ws, "A2", "FY2021–FY2025 来自 10-K；TTM = FY2025 + 9M 2026 − 9M 2025；2026E/2027E 为一致预期。")

header_row(ws, 4, ["项目", "FY2021", "FY2022", "FY2023", "FY2024", "FY2025", "TTM", "备注"])
hist = {
    "收入": [257.637, 282.836, 307.394, 350.018, 402.836, 471.098],
    "经营利润": [78.714, 74.842, 84.293, 112.390, 129.039, 156.976],
    "净利润": [76.033, 59.972, 73.795, 100.118, 132.170, 235.694],
    "经营现金流": [91.652, 91.495, 101.746, 125.299, 164.713, 179.551],
    "资本开支": [24.640, 31.485, 32.251, 52.535, 91.447, 159.804],
    "自由现金流": [67.012, 60.010, 69.495, 72.764, 73.266, 19.747],
    "稀释 EPS（美元）": [5.61, 4.56, 5.80, 8.04, 10.81, 19.15],
    "稀释股数（十亿）": [13.55, 13.159, 12.722, 12.447, 12.230, 12.324],
}
row = 5
for label, vals in hist.items():
    body_cell(ws, row, 1, label, align="left")
    for j, v in enumerate(vals, 2):
        fmt = "0.00" if "EPS" in label else ("0.000" if "股数" in label else "#,##0.0")
        body_cell(ws, row, j, v, fmt=fmt)
    row += 1

# derived rows
body_cell(ws, 13, 1, "收入增速", align="left")
for j in range(3, 8):
    body_cell(ws, 13, j, f"={get_column_letter(j)}5/{get_column_letter(j-1)}5-1", fmt="0.0%")
body_cell(ws, 14, 1, "经营利润率", align="left")
for j in range(2, 8):
    body_cell(ws, 14, j, f"={get_column_letter(j)}6/{get_column_letter(j)}5", fmt="0.0%")
body_cell(ws, 15, 1, "净利率", align="left")
for j in range(2, 8):
    body_cell(ws, 15, j, f"={get_column_letter(j)}7/{get_column_letter(j)}5", fmt="0.0%")
body_cell(ws, 16, 1, "FCF 利润率", align="left")
for j in range(2, 8):
    body_cell(ws, 16, j, f"={get_column_letter(j)}10/{get_column_letter(j)}5", fmt="0.0%")
body_cell(ws, 17, 1, "Capex / 收入", align="left")
for j in range(2, 8):
    body_cell(ws, 17, j, f"={get_column_letter(j)}9/{get_column_letter(j)}5", fmt="0.0%")

section(ws, 19, "关键读数", 8)
notes = [
    "FY2021–FY2025 收入 CAGR 约 11.8%；经营利润 CAGR 约 13.1%；EPS CAGR 约 17.8%。",
    "FY2025 经营利润率 32.0%；TTM 33.3%。Cloud 规模效应是主要增量。",
    "TTM 净利润 235.7B 含大量投资损益（SpaceX/Anthropic 等），不能直接用于估值；正常化净利润约 127B。",
    "TTM FCF 仅 19.7B，FCF/净利润 8.4%；capex 159.8B，占收入 33.9%。",
    "FY2025 回购 45.7B + 股息 10.1B；2026 年 YTD 回购为 0，转为发债/发股/优先股。",
]
for i, text in enumerate(notes, start=20):
    ws.cell(row=i, column=1, value="• " + text).font = Font(name=FONT, size=10, color=C_DARK)
    ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)
    ws.row_dimensions[i].height = 26

# --- Sheet 3: 分部 -----------------------------------------------------------
ws = wb.create_sheet("分部")
ws.sheet_view.showGridLines = False
widths(ws, {"A": 28, "B": 14, "C": 14, "D": 14, "E": 14, "F": 26})
title(ws, "A1", "Alphabet 分部：Search 现金牛 + Cloud 火箭")
subtitle(ws, "A2", "Q3 2026 数据来自 8-K 业绩稿；FY2025 来自 10-K。单位：十亿美元。")

header_row(ws, 4, ["Q3 2026", "收入", "同比", "经营利润", "经营利润率", "备注"])
q3 = [
    ("Google Search & other", 63.350, 0.120, None, None, "query 增速五年最高"),
    ("YouTube ads", 11.207, 0.092, None, None, "Shorts/订阅带动"),
    ("Google Network", 7.455, 0.014, None, None, "成熟、低增长"),
    ("Google advertising", 82.012, 0.106, None, None, "广告总盘"),
    ("Subscriptions/platforms/devices", 14.290, 0.110, None, None, "订阅增长"),
    ("Google Services", 96.302, 0.106, 39.157, 0.407, "现金牛"),
    ("Google Cloud", 30.709, 1.026, 11.121, 0.362, "backlog 561B"),
    ("Other Bets", 0.372, 0.081, -2.047, None, "亏损"),
    ("Alphabet-level", 0.0, None, -7.655, None, "AI 研发/公司费用"),
    ("Total", 127.577, 0.247, 40.576, 0.318, "收入 +25%"),
]
for i, (name, rev, growth, op, margin, note) in enumerate(q3, start=5):
    body_cell(ws, i, 1, name, align="left", bold=name in ("Google Services", "Google Cloud", "Total"))
    body_cell(ws, i, 2, rev, fmt="#,##0.0")
    body_cell(ws, i, 3, growth, fmt="0.0%")
    body_cell(ws, i, 4, op if op is not None else "—", fmt="#,##0.0" if op is not None else None)
    body_cell(ws, i, 5, margin if margin is not None else "—", fmt="0.0%" if margin is not None else None)
    body_cell(ws, i, 6, note, align="left")

section(ws, 17, "FY2025 分部", 6)
header_row(ws, 18, ["FY2025", "收入", "占比", "经营利润", "经营利润率", "备注"])
fy25 = [
    ("Google Search & other", 224.532, 0.557, None, None, ""),
    ("YouTube ads", 40.367, 0.100, None, None, ""),
    ("Google Network", 29.792, 0.074, None, None, ""),
    ("Google advertising", 294.691, 0.732, None, None, ""),
    ("Subscriptions/platforms/devices", 48.030, 0.119, None, None, ""),
    ("Google Services", 342.721, 0.851, 139.404, 0.407, ""),
    ("Google Cloud", 58.705, 0.146, 13.910, 0.237, ""),
    ("Other Bets", 1.537, 0.004, -7.515, None, ""),
    ("Hedging", -0.127, None, None, None, ""),
    ("Total", 402.836, 1.000, 129.039, 0.320, ""),
]
for i, (name, rev, share, op, margin, note) in enumerate(fy25, start=19):
    body_cell(ws, i, 1, name, align="left", bold=name in ("Google Services", "Google Cloud", "Total"))
    body_cell(ws, i, 2, rev, fmt="#,##0.0")
    body_cell(ws, i, 3, share if share is not None else "—", fmt="0.0%" if share is not None else None)
    body_cell(ws, i, 4, op if op is not None else "—", fmt="#,##0.0" if op is not None else None)
    body_cell(ws, i, 5, margin if margin is not None else "—", fmt="0.0%" if margin is not None else None)
    body_cell(ws, i, 6, note, align="left")

section(ws, 31, "Cloud/TPU 关键事实", 6)
cloud_notes = [
    "Q3 2026 Cloud 收入 30.709B，同比 +103%；经营利润 11.121B，利润率 36.2%。",
    "Cloud backlog 561B；管理层称 >55%（约 308B）将在未来 24 个月确认。",
    "TPU 8T 训练性能较上代 2.7x；TPU 8I 推理性价比高 80%；约 100 家客户使用 TPU。",
    "Gemini app 达 10 亿用户；第一方模型 API 处理约 320 亿 tokens/分钟。",
    "2026 capex 指引 195–205B；2027 将显著增加（媒体口径 >300B）；采购承诺 924B。",
]
for i, text in enumerate(cloud_notes, start=32):
    ws.cell(row=i, column=1, value="• " + text).font = Font(name=FONT, size=10, color=C_DARK)
    ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=6)
    ws.row_dimensions[i].height = 26

# --- Sheet 4: FCFF桥 ---------------------------------------------------------
ws = wb.create_sheet("FCFF桥")
ws.sheet_view.showGridLines = False
widths(ws, {"A": 34, "B": 18, "C": 34})
title(ws, "A1", "TTM FCFF 桥：为什么 FCF 只有 19.7B")
subtitle(ws, "A2", "单位：十亿美元。TTM = 2025Q4–2026Q3；经营利润已扣除 SBC。")

header_row(ws, 4, ["项目", "数值", "说明"])
rows = [
    ("经营利润（EBIT）", 156.976, "TTM"),
    ("有效税率", 0.19, "正常化假设"),
    ("NOPAT", "=B5*(1-B6)", "EBIT × (1−税率)"),
    ("+ 折旧摊销", 27.995, "TTM"),
    ("− 资本开支", 159.804, "TTM，含 AI 基础设施"),
    ("− 营运资金增加", 0.0, "正常化前"),
    ("= 正常化 FCFF（未计 WC）", "=B7+B8-B9-B10", "capex 是主要拖累"),
    ("报告 CFO", 179.551, "TTM"),
    ("− 资本开支", 159.804, "TTM"),
    ("= 报告 FCF", "=B12-B13", "TTM"),
    ("FCF / NOPAT", "=B14/B7", "现金转化率"),
    ("Capex / 收入", "=B9/471.098", "TTM"),
]
for i, (label, value, note) in enumerate(rows, start=5):
    body_cell(ws, i, 1, label, align="left", bold=label.startswith("="))
    body_cell(ws, i, 2, value, fmt="#,##0.0" if isinstance(value, (int, float)) and abs(value) > 1 else ("0.0%" if isinstance(value, float) and abs(value) <= 1 else "#,##0.0"), bold=isinstance(value, str) and value.startswith("="))
    body_cell(ws, i, 3, note, align="left")
ws["B16"].number_format = "0.0%"

section(ws, 18, "解读", 3)
notes = [
    "TTM NOPAT 127B，但 capex 160B，资本开支超过折旧摊销 132B；这是 FCF  collapse 的主因。",
    "2026 capex 指引 195–205B，2027 将显著增加；FCF 在 2026–2027 可能持续为负。",
    "TPU 库存是现金流的时点问题：签合同→备货→交付→确认收入；不是需求问题，但仍占用现金。",
    "2026 Q3 还支付了 5.2B 欧盟罚款；一次性因素放大了 FCF 压力。",
]
for i, text in enumerate(notes, start=19):
    ws.cell(row=i, column=1, value="• " + text).font = Font(name=FONT, size=10, color=C_DARK)
    ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=3)
    ws.row_dimensions[i].height = 26

# --- Sheet 5: DCF ------------------------------------------------------------
ws = wb.create_sheet("DCF")
ws.sheet_view.showGridLines = False
widths(ws, {"A": 10, "B": 10, "C": 11, "D": 10, "E": 11, "F": 11, "G": 10,
            "H": 9, "I": 10, "J": 9, "K": 10, "L": 10, "M": 9, "N": 10})

title(ws, "A1", "三情景 FCFF DCF（未来 12 个月起算，10 年明确预测）")
subtitle(ws, "A2", "单位：十亿美元，除每股数据。第 1 年约为 2026Q4–2027Q3；终值用永续增长。")

# global inputs
section(ws, 4, "全局输入", 5)
inputs = [
    ("起始收入（TTM）", TTM_REV),
    ("税率", TAX),
    ("营运资金/增量收入", NWC),
    ("净现金（扣除优先股）", NET_CASH_COMMON),
    ("股数（十亿）", SHARES),
]
for i, (label, val) in enumerate(inputs, start=5):
    body_cell(ws, i, 1, label, align="left")
    body_cell(ws, i, 2, val, fmt="#,##0.000" if isinstance(val, float) and val < 20 else "#,##0.0")
ws["A11"] = "注：净现金 = 265.098 − 125.102 − 19.000 = 120.996；优先股按清算价值扣除。"
ws["A11"].font = Font(name=FONT, size=9, color=C_GREY)
ws.merge_cells("A11:N11")

start_rows = {"bear": 14, "base": 37, "bull": 60}
for k, start in start_rows.items():
    s = SCEN[k]
    r = RESULTS[k]
    section(ws, start, f"{s['label']}情景", 14)
    body_cell(ws, start + 1, 1, "WACC", align="left")
    body_cell(ws, start + 1, 2, s["wacc"], fmt="0.0%")
    body_cell(ws, start + 1, 3, "终值增长 g", align="left")
    body_cell(ws, start + 1, 4, s["gt"], fmt="0.0%")
    header = ["Year", "收入增速", "收入", "EBIT 利润率", "EBIT", "NOPAT",
              "D&A %", "D&A", "Capex %", "Capex", "ΔNWC", "FCFF", "折现因子", "PV"]
    header_row(ws, start + 3, header)
    for t, row in enumerate(r["rows"], start=start + 4):
        body_cell(ws, t, 1, row["year"], fmt="0")
        body_cell(ws, t, 2, s["growth"][row["year"] - 1], fmt="0.0%")
        if row["year"] == 1:
            body_cell(ws, t, 3, f"=$B$5*(1+B{t})", fmt="#,##0.0")
        else:
            body_cell(ws, t, 3, f"=C{t-1}*(1+B{t})", fmt="#,##0.0")
        body_cell(ws, t, 4, s["ebit"][row["year"] - 1], fmt="0.0%")
        body_cell(ws, t, 5, f"=C{t}*D{t}", fmt="#,##0.0")
        body_cell(ws, t, 6, f"=E{t}*(1-$B$6)", fmt="#,##0.0")
        body_cell(ws, t, 7, s["da"][row["year"] - 1], fmt="0.0%")
        body_cell(ws, t, 8, f"=C{t}*G{t}", fmt="#,##0.0")
        body_cell(ws, t, 9, s["capex"][row["year"] - 1], fmt="0.0%")
        body_cell(ws, t, 10, f"=C{t}*I{t}", fmt="#,##0.0")
        body_cell(ws, t, 11, f"=(C{t}-$B$5)*$B$7" if row["year"] == 1 else f"=(C{t}-C{t-1})*$B$7", fmt="#,##0.0")
        body_cell(ws, t, 12, f"=F{t}+H{t}-J{t}-K{t}", fmt="#,##0.0")
        body_cell(ws, t, 13, f"=1/(1+$B${start+1})^{t-start-3}", fmt="0.000")
        body_cell(ws, t, 14, f"=L{t}*M{t}", fmt="#,##0.0")
    end = start + 13
    body_cell(ws, end + 1, 1, "终值", align="left", bold=True)
    body_cell(ws, end + 1, 2, f"=L{end}*(1+$D${start+1})/($B${start+1}-$D${start+1})", fmt="#,##0.0")
    body_cell(ws, end + 2, 1, "PV 终值", align="left")
    body_cell(ws, end + 2, 2, f"=B{end+1}/(1+$B${start+1})^10", fmt="#,##0.0")
    body_cell(ws, end + 3, 1, "EV", align="left", bold=True)
    body_cell(ws, end + 3, 2, f"=SUM(N{start+4}:N{end})+B{end+2}", fmt="#,##0.0")
    body_cell(ws, end + 4, 1, "净现金", align="left")
    body_cell(ws, end + 4, 2, f"=$B$8", fmt="#,##0.0")
    body_cell(ws, end + 5, 1, "股权价值", align="left", bold=True)
    body_cell(ws, end + 5, 2, f"=B{end+3}+B{end+4}", fmt="#,##0.0")
    body_cell(ws, end + 6, 1, "每股价值", align="left", bold=True)
    body_cell(ws, end + 6, 2, f"=B{end+5}/$B$9", fmt="$#,##0.00")
    body_cell(ws, end + 6, 3, f"TV 占比 {r['tv_share']:.0%}；终值 FCF {r['term_fcf']:.0f}B；终值收入 {r['term_rev']:.0f}B", align="left")

# Base-case FCFF line chart
chart = LineChart()
chart.title = "基准情景 FCFF 路径（十亿美元）"
chart.y_axis.title = "FCFF"
chart.x_axis.title = "Year"
data = Reference(ws, min_col=12, min_row=40, max_row=50)
cats = Reference(ws, min_col=1, min_row=41, max_row=50)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.height = 7
chart.width = 13
ws.add_chart(chart, "P14")

# --- Sheet 6: 敏感性 ---------------------------------------------------------
ws = wb.create_sheet("敏感性")
ws.sheet_view.showGridLines = False
widths(ws, {"A": 24, "B": 14, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14})
title(ws, "A1", "估值敏感性：WACC、终值 capex、终值利润率")
subtitle(ws, "A2", "基准情景；每股价值（美元）。由 Python 独立计算，Excel 表为静态值。")

header_row(ws, 4, ["WACC", "每股价值"])
for i, (w, ps) in enumerate(zip(WACC_GRID, WACC_PS), start=5):
    body_cell(ws, i, 1, w, fmt="0.0%")
    body_cell(ws, i, 2, r0(ps), fmt="$#,##0")

header_row(ws, 14, ["终值 Capex / 收入", "每股价值"])
for i, (c, ps) in enumerate(zip(CAPEX_GRID, CAPEX_PS), start=15):
    body_cell(ws, i, 1, c, fmt="0.0%")
    body_cell(ws, i, 2, r0(ps), fmt="$#,##0")

header_row(ws, 22, ["终值 EBIT 利润率", "每股价值"])
for i, (m, ps) in enumerate(zip(MARGIN_GRID, MARGIN_PS), start=23):
    body_cell(ws, i, 1, m, fmt="0.0%")
    body_cell(ws, i, 2, r0(ps), fmt="$#,##0")

section(ws, 31, "反推：当前价格隐含的 10 年收入 CAGR", 3)
body_cell(ws, 32, 1, "盘后价 $291.50", align="left")
body_cell(ws, 32, 2, IMPLIED_G_AH, fmt="0.0%")
body_cell(ws, 33, 1, "收盘价 $324.89", align="left")
body_cell(ws, 33, 2, IMPLIED_G_CLOSE, fmt="0.0%")
body_cell(ws, 34, 1, "基准情景 10 年 CAGR", align="left")
body_cell(ws, 34, 2, (RESULTS["base"]["term_rev"] / TTM_REV) ** (1 / 10) - 1, fmt="0.0%")

section(ws, 36, "SOTP 交叉检查（每股美元）", 7)
header_row(ws, 37, ["Services EBIT 倍数", "Cloud 25x", "Cloud 30x", "Cloud 35x", "", "", ""])
row = 38
for sm in [17, 18, 19, 20]:
    body_cell(ws, row, 1, f"{sm}x", align="left")
    for j, cm in enumerate([25, 30, 35], start=2):
        ps = next(x[3] for x in SOTP if x[0] == sm and x[1] == cm)
        body_cell(ws, row, j, r0(ps), fmt="$#,##0")
    row += 1
body_cell(ws, 43, 1, "注：Services/Cloud 用 TTM 经营利润；公司费用按 15x 扣除；净现金 + 非上市证券 50% 计入。", align="left")
ws.merge_cells("A43:G43")
ws["A43"].font = Font(name=FONT, size=9, color=C_GREY)

# --- Sheet 7: 来源 -----------------------------------------------------------
ws = wb.create_sheet("来源")
ws.sheet_view.showGridLines = False
widths(ws, {"A": 34, "B": 50, "C": 16, "D": 50})
title(ws, "A1", "数据来源与日期")
subtitle(ws, "A2", "研究截止 2026-10-29；价格、一致预期、目标价均为当日或前一交易日数据。")
header_row(ws, 4, ["数据", "来源", "日期", "链接/备注"])
sources = [
    ("FY2025 10-K", "SEC EDGAR, Alphabet Inc.", "2026-02-05", "https://www.sec.gov/Archives/edgar/data/1652044/000165204426000018/goog-20251231.htm"),
    ("Q3 2026 业绩", "SEC 8-K Exhibit 99.1", "2026-10-28", "https://www.sec.gov/Archives/edgar/data/1652044/000165204426000080/googexhibit991q32026.htm"),
    ("Q3 2026 电话会", "Alphabet earnings call / 媒体整理", "2026-10-28", "Cloud backlog 561B；2027 capex 显著增加；924B 采购承诺"),
    ("价格/市值", "Yahoo Finance / StockAnalysis", "2026-10-28", "收盘 324.89；盘后约 291.50"),
    ("一致预期", "StockAnalysis / Investing.com / TipRanks", "2026-10", "2026E 收入 ~499B；2027E ~620B；调整后 EPS 11.89/15.19"),
    ("目标价", "S&P Global / TipRanks / Investing.com", "2026-10-28", "均值约 430 美元（财报前，可能已过时）"),
    ("10Y 美债", "YCharts / CNBC", "2026-10-28", "5.03%"),
    ("反垄断", "TechCrunch / CRS / 媒体", "2026-09", "Chrome 未被拆分；行为救济；上诉中"),
]
for i, (a, b, c, d) in enumerate(sources, start=5):
    body_cell(ws, i, 1, a, align="left")
    body_cell(ws, i, 2, b, align="left")
    body_cell(ws, i, 3, c, align="center")
    body_cell(ws, i, 4, d, align="left")

wb.save(XLSX)
print("saved:", XLSX)
print("scenario:", {k: round(v["ps"], 2) for k, v in RESULTS.items()})
print("IRR @291.5:", {k: round(v * 100, 2) for k, v in IRR_AH.items()})
print("IRR @324.89:", {k: round(v * 100, 2) for k, v in IRR_CLOSE.items()})
print("implied CAGR:", round(IMPLIED_G_AH * 100, 2), round(IMPLIED_G_CLOSE * 100, 2))
