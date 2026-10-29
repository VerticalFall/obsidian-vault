# -*- coding: utf-8 -*-
"""Build Microsoft (MSFT) Buffett-framework valuation workbook.

All figures in USD billions unless noted. Data cut-off: 2026-10-29.
Official financials: FY2026 10-K (ended 2026-06-30, filed 2026-07-29) and
Q1 FY2027 10-Q / 8-K earnings release (quarter ended 2026-09-30, filed
2026-10-28). Reference price: $546.39 after-hours on 2026-10-28; last
regular close $531.01.
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
XLSX = BASE / "MSFT-估值模型-2026-10-29.xlsx"

# ----------------------------------------------------------------------------
# Inputs and computed outputs
# ----------------------------------------------------------------------------
PRICE_AH = 546.39          # 2026-10-28 after-hours indicative
PRICE_CLOSE = 531.01       # 2026-10-28 regular close
SHARES = 7.427             # billion, 2026-09-30
CASH = 64.509              # cash + short-term investments
DEBT = 35.481              # current + long-term debt
FIN_LEASE = 70.973         # finance lease liabilities, Sep-30-2026 (debt-like)
OP_LEASE = 27.477          # operating lease liabilities, Sep-30-2026 (memo)
EQUITY_INV = 37.808        # equity and other investments (incl. OpenAI book value)
NET_DEBT = DEBT + FIN_LEASE - CASH   # positive = net debt
MCAP_AH = PRICE_AH * SHARES
MCAP_CLOSE = PRICE_CLOSE * SHARES
EV_AH = MCAP_AH + NET_DEBT
EV_CLOSE = MCAP_CLOSE + NET_DEBT

TTM_REV = 346.445
TTM_EBIT = 163.191
TTM_NOPAT = TTM_EBIT * 0.81
TTM_FCF = 59.957
TTM_EPS_GAAP = 19.15
TTM_EPS_ADJ = 18.04
REV_2027E = 406.0
REV_2028E = 471.0
EPS_2027E = 20.00
EPS_2028E = 23.69
DIV_PS = 3.64
TAX = 0.19
NWC = 0.02

SCEN = {
    "bear": dict(
        label="悲观", wacc=0.093, gt=0.030,
        growth=[0.14, 0.12, 0.09, 0.07, 0.05, 0.04, 0.03, 0.03, 0.02, 0.02],
        ebit=[0.455, 0.440, 0.430, 0.425, 0.425, 0.425, 0.425, 0.425, 0.425, 0.425],
        da=[0.135, 0.155, 0.165, 0.165, 0.155, 0.145, 0.135, 0.130, 0.125, 0.120],
        capex=[0.375, 0.360, 0.330, 0.300, 0.280, 0.260, 0.240, 0.220, 0.200, 0.185],
    ),
    "base": dict(
        label="基准", wacc=0.088, gt=0.035,
        growth=[0.172, 0.16, 0.14, 0.12, 0.10, 0.09, 0.08, 0.07, 0.06, 0.05],
        ebit=[0.470, 0.465, 0.465, 0.470, 0.475, 0.480, 0.480, 0.480, 0.480, 0.480],
        da=[0.130, 0.150, 0.160, 0.160, 0.150, 0.140, 0.130, 0.125, 0.120, 0.115],
        capex=[0.365, 0.320, 0.270, 0.230, 0.200, 0.180, 0.165, 0.155, 0.150, 0.145],
    ),
    "bull": dict(
        label="乐观", wacc=0.082, gt=0.040,
        growth=[0.19, 0.18, 0.16, 0.14, 0.12, 0.10, 0.09, 0.08, 0.07, 0.06],
        ebit=[0.480, 0.485, 0.490, 0.495, 0.495, 0.495, 0.495, 0.495, 0.495, 0.495],
        da=[0.130, 0.150, 0.160, 0.160, 0.150, 0.140, 0.130, 0.125, 0.120, 0.115],
        capex=[0.355, 0.305, 0.250, 0.210, 0.180, 0.160, 0.145, 0.135, 0.128, 0.122],
    ),
}


def dcf(s):
    rev = TTM_REV
    prev = rev
    pv = 0.0
    rows = []
    for t in range(10):
        rev = prev * (1 + s["growth"][t])
        ebit = rev * s["ebit"][t]
        nopat = ebit * (1 - TAX)
        da = rev * s["da"][t]
        capex = rev * s["capex"][t]
        dnwc = (rev - prev) * NWC
        fcf = nopat + da - capex - dnwc
        df = 1 / (1 + s["wacc"]) ** (t + 1)
        rows.append(dict(year=t + 1, growth=s["growth"][t], rev=rev, ebit=ebit,
                         nopat=nopat, da=da, capex=capex, dnwc=dnwc,
                         fcf=fcf, df=df, pv=fcf * df))
        pv += fcf * df
        prev = rev
    tv = rows[-1]["fcf"] * (1 + s["gt"]) / (s["wacc"] - s["gt"])
    pvtv = tv / (1 + s["wacc"]) ** 10
    ev = pv + pvtv
    equity = ev - NET_DEBT
    return dict(rows=rows, pv=pv, tv=tv, pvtv=pvtv, ev=ev,
                equity=equity, ps=equity / SHARES,
                tv_share=pvtv / ev, term_fcf=rows[-1]["fcf"],
                term_rev=rows[-1]["rev"])


RESULTS = {k: dcf(v) for k, v in SCEN.items()}


def bisect(f, lo, hi, tol=1e-10):
    flo = f(lo)
    for _ in range(300):
        mid = (lo + hi) / 2
        fm = f(mid)
        if abs(fm) < tol:
            return mid
        if flo * fm < 0:
            hi = mid
        else:
            lo = mid
            flo = fm
    return (lo + hi) / 2


def r0(x):
    """Round half-up for display tables (avoid Python banker's rounding)."""
    return int(x + 0.5 + 1e-6) if x >= 0 else -int(-x + 0.5 + 1e-6)


def irr_at(price, s, rows, term_fcf):
    def npv(r):
        p = sum(row["fcf"] / (1 + r) ** row["year"] for row in rows)
        tv = term_fcf * (1 + s["gt"]) / (r - s["gt"])
        return p + tv / (1 + r) ** 10 - NET_DEBT - price * SHARES
    return bisect(npv, s["gt"] + 1e-4, 0.50)


IRR_AH = {k: irr_at(PRICE_AH, v, RESULTS[k]["rows"], RESULTS[k]["term_fcf"]) for k, v in SCEN.items()}
IRR_CLOSE = {k: irr_at(PRICE_CLOSE, v, RESULTS[k]["rows"], RESULTS[k]["term_fcf"]) for k, v in SCEN.items()}

# WACC sensitivity on base case
WACC_GRID = [0.080, 0.083, 0.086, 0.088, 0.090, 0.093, 0.096, 0.100]
WACC_PS = []
for w in WACC_GRID:
    s = dict(SCEN["base"])
    s["wacc"] = w
    WACC_PS.append(dcf(s)["ps"])

# Terminal capex sensitivity on base case
CAPEX_GRID = [0.130, 0.140, 0.150, 0.160, 0.170, 0.180]
CAPEX_PS = []
for c in CAPEX_GRID:
    s = dict(SCEN["base"])
    s["capex"] = list(SCEN["base"]["capex"][:-1]) + [c]
    CAPEX_PS.append(dcf(s)["ps"])

# Terminal EBIT margin sensitivity on base case
MARGIN_GRID = [0.440, 0.450, 0.460, 0.470, 0.480, 0.490, 0.500]
MARGIN_PS = []
for m in MARGIN_GRID:
    s = dict(SCEN["base"])
    s["ebit"] = list(SCEN["base"]["ebit"][:-1]) + [m]
    MARGIN_PS.append(dcf(s)["ps"])


def value_const_g(g, base=SCEN["base"]):
    rev = TTM_REV
    prev = rev
    pv = 0.0
    last = 0.0
    for t in range(10):
        rev = prev * (1 + g)
        nopat = rev * base["ebit"][t] * (1 - TAX)
        fcf = nopat + rev * base["da"][t] - rev * base["capex"][t] - (rev - prev) * NWC
        pv += fcf / (1 + base["wacc"]) ** (t + 1)
        last = fcf
        prev = rev
    tv = last * (1 + base["gt"]) / (base["wacc"] - base["gt"])
    return pv + tv / (1 + base["wacc"]) ** 10 - NET_DEBT


IMPLIED_G_AH = bisect(lambda g: value_const_g(g) - PRICE_AH * SHARES, -0.2, 0.4)
IMPLIED_G_CLOSE = bisect(lambda g: value_const_g(g) - PRICE_CLOSE * SHARES, -0.2, 0.4)

# SOTP cross-check. Segment EBIT split uses Q1 FY27 mix applied to TTM EBIT:
# Agents & Infra 40.558 / 45.915 = 88.33%; Devices & Consumer 11.67%.
AI_EBIT = TTM_EBIT * 40.558 / 45.915
DC_EBIT = TTM_EBIT * 5.357 / 45.915
SOTP = []
for sm in [25, 28, 30, 32]:
    for dm in [10.0, 11.5, 13.0]:
        eq = AI_EBIT * sm + DC_EBIT * dm + EQUITY_INV - NET_DEBT
        SOTP.append((sm, dm, eq, eq / SHARES))

# ----------------------------------------------------------------------------
# Styling helpers
# ----------------------------------------------------------------------------
FONT = "Microsoft YaHei"
C_DARK = "1F2933"
C_LIGHT = "F7F4EC"
C_CREAM = "E8E2D5"
C_RED = "B23A2E"
C_GREEN = "2F6B4F"
C_GOLD = "C9A227"
C_BLUE = "1A73E8"
C_GREY = "5B6770"
C_YELLOW = "FFF2CC"
C_WHITE = "FFFFFF"

thin = Side(style="thin", color="D8D2C4")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def title(ws, cell, text, size=16):
    ws[cell] = text
    ws[cell].font = Font(name=FONT, size=size, bold=True, color=C_DARK)


def subtitle(ws, cell, text):
    ws[cell] = text
    ws[cell].font = Font(name=FONT, size=10, color=C_GREY)


def section(ws, row, text, ncols=8):
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=FONT, size=11, bold=True, color=C_DARK)
    for c in range(1, ncols + 1):
        ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor=C_CREAM)
        ws.cell(row=row, column=c).border = BORDER


def header_row(ws, row, values):
    for j, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=j, value=v)
        cell.font = Font(name=FONT, size=9, bold=True, color=C_WHITE)
        cell.fill = PatternFill("solid", fgColor=C_DARK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def body_cell(ws, row, col, value, fmt=None, bold=False, align="right", color=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(name=FONT, size=9, bold=bold, color=color or C_DARK)
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.border = BORDER
    if fmt:
        cell.number_format = fmt
    return cell


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


# ----------------------------------------------------------------------------
# Workbook
# ----------------------------------------------------------------------------
wb = Workbook()

# --- Sheet 1: 结论 -----------------------------------------------------------
ws = wb.active
ws.title = "结论"
ws.sheet_view.showGridLines = False
widths(ws, {"A": 30, "B": 17, "C": 20, "D": 22, "E": 14})

title(ws, "A1", "微软（MSFT）巴菲特框架估值")
subtitle(ws, "A2", "研究日 2026-10-29｜盘后参考价 $546.39｜10/28 收盘 $531.01｜最新财报：FY2027 Q1（2026-09-30）")
subtitle(ws, "A3", "FY2026 10-K 与 FY2027 Q1 10-Q 均已提交；市值按 7.427B 股计算。")

section(ws, 5, "一句话结论", 5)
ws["A6"] = "好生意，合理价格，但没有安全边际。Azure 加速证明 AI 需求是真的，但 FY2027 capex 约 170B 把 FCF 压在 60B 附近。结论：观察，等 FCF 拐点。"
ws["A6"].font = Font(name=FONT, size=11, bold=True, color=C_RED)
ws["A6"].alignment = Alignment(wrap_text=True, vertical="center")
ws.merge_cells("A6:E6")
ws.row_dimensions[6].height = 44

section(ws, 8, "关键数据（十亿美元 / 倍 / 百分比）", 5)
header_row(ws, 9, ["项目", "数值", "备注"])
rows = [
    ("盘后参考价（美元）", PRICE_AH, "2026-10-28 盘后，约 +2.9%", "$#,##0.00"),
    ("10/28 收盘价（美元）", PRICE_CLOSE, "财报前收盘", "$#,##0.00"),
    ("期末股数（十亿股）", SHARES, "2026-09-30", "0.000"),
    ("市值 @ 盘后（十亿美元）", "=B10*B12", "价格 × 股数", "#,##0.0"),
    ("净负债（含融资租赁）", round(NET_DEBT, 3), "债务 35.5 + 融资租赁 71.0 − 现金 64.5", "#,##0.0"),
    ("企业价值 @ 盘后", "=B13+B14", "市值 + 净负债", "#,##0.0"),
    ("TTM 收入", TTM_REV, "FY2026Q2–FY2027Q1", "#,##0.0"),
    ("TTM 经营利润", TTM_EBIT, "经营利润，剔除 OI&E", "#,##0.0"),
    ("TTM NOPAT", round(TTM_NOPAT, 3), "经营利润 × (1−19%)", "#,##0.0"),
    ("TTM 自由现金流", TTM_FCF, "CFO − 现金 capex", "#,##0.0"),
    ("NTM 收入（2026-10 至 2027-09）", REV_2027E, "TTM × (1+17.2%)，DCF 第 1 年", "#,##0.0"),
    ("DCF 第 2 年收入", REV_2028E, "NTM × (1+16%)", "#,##0.0"),
    ("FY2027E 调整后 EPS（美元）", EPS_2027E, "本文估计，非一致预期", "0.00"),
    ("FY2028E 调整后 EPS（美元）", EPS_2028E, "一致预期 23.69", "0.00"),
    ("EV / TTM EBIT", "=B15/B17", "盘后价", '0.0"x"'),
    ("P/E FY2027E", "=B10/B22", "盘后价 / 调整后 EPS", '0.0"x"'),
    ("P/E FY2028E", "=B10/B23", "盘后价 / 调整后 EPS", '0.0"x"'),
    ("FCF 收益率", "=B19/B13", "TTM FCF / 市值", "0.00%"),
    ("股息率", "=3.64/B10", "年化股息 $3.64", "0.00%"),
]
for i, (label, value, note, fmt) in enumerate(rows, start=10):
    body_cell(ws, i, 1, label, align="left")
    body_cell(ws, i, 2, value, fmt=fmt, bold=isinstance(value, str) and value.startswith("="))
    body_cell(ws, i, 3, note, align="left")

section(ws, 31, "三情景 DCF（FCFF，WACC 折现，减净负债）", 5)
header_row(ws, 32, ["情景", "每股价值（美元）", "较盘后价", "安全边际（1−价/值）", "10 年 IRR"])
for i, k in enumerate(["bear", "base", "bull"], start=33):
    r = RESULTS[k]
    dcf_row = {"bear": 32, "base": 55, "bull": 78}[k]
    body_cell(ws, i, 1, SCEN[k]["label"], align="left", bold=(k == "base"))
    body_cell(ws, i, 2, f"='DCF'!B{dcf_row}", fmt="$#,##0.00", bold=(k == "base"))
    body_cell(ws, i, 3, f"=B{i}/$B$10-1", fmt="0.0%")
    body_cell(ws, i, 4, f"=1-$B$10/B{i}", fmt="0.0%")
    body_cell(ws, i, 5, round(IRR_AH[k], 4), fmt="0.00%")
body_cell(ws, 36, 1, "盘后价", align="left", bold=True, color=C_RED)
body_cell(ws, 36, 2, PRICE_AH, fmt="$#,##0.00", bold=True, color=C_RED)
body_cell(ws, 36, 3, 0, fmt="0.0%")
body_cell(ws, 36, 4, 0, fmt="0.0%")
body_cell(ws, 36, 5, "—", align="center")

section(ws, 38, "结论与观察条件", 5)
notes = [
    "基准价值约 540 美元，与盘后价 546.39 几乎持平（高 1%）；收盘价 531.01 低 2%。",
    "悲观 244；基准 540；乐观 839；概率加权（25/50/25）约 540.8。",
    "反推：盘后价隐含未来 10 年收入 CAGR 约 10.9%，基准 10.35%，只差 0.5pp。",
    "若要求 20% 安全边际：基准观察价约 433 美元；按 CAPM 9.3% 的基准 489 美元，观察价约 391 美元。",
    "结论是“观察”，不是“卖出”。已持有不必因一个季度超预期追高；未持有可等 FCF 拐点或回调。",
    "这不是投资建议；估值对 capex 正常化速度、Azure 增速和 WACC 高度敏感。",
]
for i, text in enumerate(notes, start=39):
    ws.cell(row=i, column=1, value="• " + text).font = Font(name=FONT, size=10, color=C_DARK)
    ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=5)
    ws.row_dimensions[i].height = 24

chart = BarChart()
chart.type = "bar"
chart.style = 10
chart.title = "每股价值：三情景 vs 盘后价（美元）"
chart.y_axis.title = "美元/股"
data = Reference(ws, min_col=2, min_row=32, max_row=36)
cats = Reference(ws, min_col=1, min_row=33, max_row=36)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.height = 7.5
chart.width = 13
chart.legend = None
ws.add_chart(chart, "G9")

# --- Sheet 2: 历史财务 -------------------------------------------------------
ws = wb.create_sheet("历史财务")
ws.sheet_view.showGridLines = False
widths(ws, {"A": 28, "B": 11, "C": 11, "D": 11, "E": 11, "F": 11, "G": 12, "H": 34})
title(ws, "A1", "微软五年财务与现金流（十亿美元，除每股数据）")
subtitle(ws, "A2", "FY2022–FY2026 来自 10-K；TTM = FY2026 + FY2027Q1 − FY2026Q1。")

header_row(ws, 4, ["项目", "FY2022", "FY2023", "FY2024", "FY2025", "FY2026", "TTM", "备注"])
hist = {
    "收入": [198.270, 211.915, 245.122, 281.724, 331.839, 346.445],
    "经营利润": [83.383, 88.523, 109.433, 128.528, 155.237, 163.191],
    "净利润": [72.738, 72.361, 88.136, 101.832, 133.749, 142.631],
    "经营现金流": [89.035, 87.582, 118.548, 136.162, 182.935, 195.878],
    "资本开支（现金）": [23.886, 28.107, 44.477, 64.551, 115.948, 135.921],
    "自由现金流": [65.149, 59.475, 74.071, 71.611, 66.987, 59.957],
    "稀释 EPS（美元）": [9.65, 9.68, 11.80, 13.64, 17.95, 19.15],
    "稀释股数（十亿）": [7.540, 7.472, 7.469, 7.465, 7.453, 7.440],
}
row = 5
for label, vals in hist.items():
    body_cell(ws, row, 1, label, align="left")
    for j, v in enumerate(vals, 2):
        fmt = "0.00" if "EPS" in label else ("0.000" if "股数" in label else "#,##0.0")
        body_cell(ws, row, j, v, fmt=fmt)
    row += 1

body_cell(ws, 13, 1, "收入增速", align="left")
for j in range(3, 8):
    body_cell(ws, 13, j, f"={get_column_letter(j)}5/{get_column_letter(j-1)}5-1", fmt="0.0%")
body_cell(ws, 14, 1, "经营利润率", align="left")
for j in range(2, 8):
    body_cell(ws, 14, j, f"={get_column_letter(j)}6/{get_column_letter(j)}5", fmt="0.0%")
body_cell(ws, 15, 1, "净利率", align="left")
for j in range(2, 8):
    body_cell(ws, 15, j, f"={get_column_letter(j)}7/{get_column_letter(j)}5", fmt="0.0%")
body_cell(ws, 16, 1, "FCF 利润率", align="left")
for j in range(2, 8):
    body_cell(ws, 16, j, f"={get_column_letter(j)}10/{get_column_letter(j)}5", fmt="0.0%")
body_cell(ws, 17, 1, "Capex / 收入", align="left")
for j in range(2, 8):
    body_cell(ws, 17, j, f"={get_column_letter(j)}9/{get_column_letter(j)}5", fmt="0.0%")
body_cell(ws, 18, 1, "FCF / NOPAT", align="left")
for j in range(2, 8):
    body_cell(ws, 18, j, f"={get_column_letter(j)}10/({get_column_letter(j)}6*0.81)", fmt="0.0%")

section(ws, 20, "关键读数", 8)
notes = [
    "FY2022–FY2026 收入 CAGR 约 13.8%；经营利润 CAGR 约 16.8%；EPS CAGR 约 16.8%。",
    "FY2026 收入 +17.8%，经营利润 +20.8%，GAAP EPS +31.6%；调整后 EPS +22.3%（剔除 OpenAI 投资损益）。",
    "FY2026 经营利润率 46.8%；TTM 47.1%。Q1 FY2027 升至 49.8%，主因收入超预期与费用纪律。",
    "FY2026 FCF 67.0B，同比 -6.5%；capex 115.9B，同比 +80%，占收入 34.9%。",
    "TTM FCF 60.0B，FCF/NOPAT 45%；capex 135.9B，capex/D&A 3.2x。",
    "FY2026 现金口径回购 22.3B（含员工税 5.6B；公开市场 16.7B）+ 股息 26.4B = 48.7B，占 FCF 73%；截至 2026-09-30 回购授权剩余 37.7B。",
]
for i, text in enumerate(notes, start=21):
    ws.cell(row=i, column=1, value="• " + text).font = Font(name=FONT, size=10, color=C_DARK)
    ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)
    ws.row_dimensions[i].height = 26

# --- Sheet 3: 分部 -----------------------------------------------------------
ws = wb.create_sheet("分部")
ws.sheet_view.showGridLines = False
widths(ws, {"A": 30, "B": 14, "C": 14, "D": 14, "E": 14, "F": 34})
title(ws, "A1", "微软分部：Agents & Infra 是利润引擎，Devices & Consumer 在收缩")
subtitle(ws, "A2", "Q1 FY2027 新分部口径（2026-09 重组）；FY2026 为旧分部口径，不可直接比较。单位：十亿美元。")

header_row(ws, 4, ["Q1 FY2027 分部", "收入", "同比", "经营利润", "经营利润率", "备注"])
seg = [
    ("Agents and Infra", 76.795, 0.245, 40.558, 0.528, "Azure/M365 cloud/LinkedIn/Dynamics"),
    ("Devices and Consumer", 15.484, -0.032, 5.357, 0.346, "Windows/Xbox/Search/consumer"),
    ("Total", 92.279, 0.188, 45.915, 0.498, "收入 +19%，经营利润 +21%"),
]
for i, (name, rev, growth, op, margin, note) in enumerate(seg, start=5):
    body_cell(ws, i, 1, name, align="left", bold=name == "Total")
    body_cell(ws, i, 2, rev, fmt="#,##0.0")
    body_cell(ws, i, 3, growth, fmt="0.0%")
    body_cell(ws, i, 4, op, fmt="#,##0.0")
    body_cell(ws, i, 5, margin, fmt="0.0%")
    body_cell(ws, i, 6, note, align="left")

section(ws, 9, "Q1 FY2027 产品/服务收入", 6)
header_row(ws, 10, ["产品/服务", "收入", "同比", "", "", "备注"])
prod = [
    ("Azure", 33.165, 0.482, "首次披露绝对收入；增速 48%"),
    ("Microsoft 365 cloud", 27.778, 0.172, "商业云 +17%"),
    ("Productivity/server licensing", 8.407, -0.031, "传统许可承压"),
    ("Search and advertising", 6.554, 0.103, "ex-TAC +12%"),
    ("XBOX", 5.379, -0.023, "内容与服务 -2%"),
    ("Industry solutions", 5.337, 0.091, "行业云 +10%"),
    ("Windows OEM and devices", 3.551, -0.220, "OEM -22%"),
    ("Frontier and support", 2.108, 0.043, "支持服务"),
    ("Total", 92.279, 0.188, ""),
]
for i, (name, rev, growth, note) in enumerate(prod, start=11):
    body_cell(ws, i, 1, name, align="left", bold=name == "Total")
    body_cell(ws, i, 2, rev, fmt="#,##0.0")
    body_cell(ws, i, 3, growth, fmt="0.0%")
    body_cell(ws, i, 6, note, align="left")

section(ws, 21, "FY2026 旧分部（仅供历史对比）", 6)
header_row(ws, 22, ["FY2026", "收入", "同比", "经营利润", "经营利润率", "备注"])
old = [
    ("Productivity and Business Processes", 139.996, 0.16, 83.879, 0.599, "M365/LinkedIn/Dynamics"),
    ("Intelligent Cloud", 137.791, 0.30, 56.972, 0.413, "Azure/服务器"),
    ("More Personal Computing", 54.052, -0.01, 14.386, 0.266, "Windows/Xbox/Search"),
    ("Total", 331.839, 0.178, 155.237, 0.468, ""),
]
for i, (name, rev, growth, op, margin, note) in enumerate(old, start=23):
    body_cell(ws, i, 1, name, align="left", bold=name == "Total")
    body_cell(ws, i, 2, rev, fmt="#,##0.0")
    body_cell(ws, i, 3, growth, fmt="0.0%")
    body_cell(ws, i, 4, op, fmt="#,##0.0")
    body_cell(ws, i, 5, margin, fmt="0.0%")
    body_cell(ws, i, 6, note, align="left")

section(ws, 29, "Cloud / AI 关键事实", 6)
cloud_notes = [
    "Microsoft Cloud FY2026 收入 214.4B，+27%，占收入 64.6%；Q1 FY2027 63.9B，+30%。",
    "Azure FY2026 +41%；Q1 FY2027 +48%，绝对收入 33.165B；Q2 指引 49–50% cc。",
    "商业 RPO 677B，+73%；FY2026 678B，+84%，加权平均期限约 2.3 年。",
    "Microsoft Cloud 毛利率 FY2026 66%，低于 FY2025 的 69%；Q1 FY2027 66%，低于上年同期 68%。",
    "Copilot 付费席位 30M+（2026-06），约占 450M+ M365 商业席位的 7%；目标 2027 年中 60M。",
    "FY2026 capex 115.9B，+80%；FY2027 指引约 170B，Q1 45.2B，Q2 >50B，下半年加速。",
    "OpenAI：约 27% 股权（2025-10 重组），当时估值 135B；IP 权利至 2032；20% 收入分成；250B Azure 采购承诺。",
]
for i, text in enumerate(cloud_notes, start=30):
    ws.cell(row=i, column=1, value="• " + text).font = Font(name=FONT, size=10, color=C_DARK)
    ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=6)
    ws.row_dimensions[i].height = 28

# --- Sheet 4: FCFF桥 ---------------------------------------------------------
ws = wb.create_sheet("FCFF桥")
ws.sheet_view.showGridLines = False
widths(ws, {"A": 34, "B": 18, "C": 46})
title(ws, "A1", "TTM FCFF 桥：为什么 132B 的 NOPAT 只变成 60B 的 FCF")
subtitle(ws, "A2", "单位：十亿美元。TTM = FY2026Q2–FY2027Q1；经营利润已扣除 SBC。")

header_row(ws, 4, ["项目", "数值", "说明"])
rows = [
    ("经营利润（EBIT）", 163.191, "TTM"),
    ("有效税率", 0.19, "正常化假设"),
    ("NOPAT", "=B5*(1-B6)", "EBIT × (1−税率)"),
    ("+ 折旧摊销", 42.131, "TTM"),
    ("+ 股权激励（SBC）", 12.556, "非现金，但会稀释"),
    ("− 资本开支（现金）", 135.921, "TTM，含 AI 基础设施"),
    ("+ 营运资金及其他", 9.006, "RPO/预收款带来的现金流入"),
    ("= 自由现金流（FCFF 口径）", "=B7+B8+B9-B10+B11", "与报告 FCF 对齐"),
    ("报告 CFO", 195.878, "TTM"),
    ("− 资本开支（现金）", 135.921, "TTM"),
    ("= 报告 FCF", "=B13-B14", "TTM"),
    ("FCF / NOPAT", "=B15/B7", "现金转化率"),
    ("Capex / 收入", "=B10/346.445", "TTM"),
    ("Capex / D&A", "=B10/B8", "资本开支是折旧的 3.2 倍"),
]
for i, (label, value, note) in enumerate(rows, start=5):
    body_cell(ws, i, 1, label, align="left", bold=label.startswith("="))
    fmt = ("0.0%" if isinstance(value, float) and abs(value) <= 1
           else "#,##0.0")
    body_cell(ws, i, 2, value, fmt=fmt, bold=isinstance(value, str) and value.startswith("="))
    body_cell(ws, i, 3, note, align="left")
ws["B16"].number_format = "0.0%"
ws["B17"].number_format = "0.0%"
ws["B18"].number_format = '0.0"x"'

section(ws, 20, "解读", 3)
notes = [
    "TTM NOPAT 132B，但 capex 136B，资本开支已经超过全部 NOPAT；这是 FCF 只有 60B 的主因。",
    "D&A 42B，capex 136B，差额 94B；这些新增产能会在未来 3–5 年转化为折旧，压制利润率。",
    "FY2027 capex 指引约 170B，是 FY2026 的 1.47 倍；公司仍预计全年 FCF 为正。",
    "营运资金贡献 +9B，主要来自 RPO/预收款；这部分不是可持续的自由现金流，估值时应打折。",
    "SBC 12.6B 加回后 FCF 才与报告口径一致；公开市场回购 16.7B > SBC，净稀释为负。",
]
for i, text in enumerate(notes, start=21):
    ws.cell(row=i, column=1, value="• " + text).font = Font(name=FONT, size=10, color=C_DARK)
    ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=3)
    ws.row_dimensions[i].height = 28

# --- Sheet 5: DCF ------------------------------------------------------------
ws = wb.create_sheet("DCF")
ws.sheet_view.showGridLines = False
widths(ws, {"A": 10, "B": 10, "C": 11, "D": 10, "E": 11, "F": 11, "G": 10,
            "H": 9, "I": 10, "J": 9, "K": 10, "L": 10, "M": 9, "N": 10})

title(ws, "A1", "三情景 FCFF DCF（未来 12 个月起算，10 年明确预测）")
subtitle(ws, "A2", "单位：十亿美元，除每股数据。第 1 年约为 FY2027Q2–FY2028Q1；终值用永续增长。")

section(ws, 4, "全局输入", 5)
inputs = [
    ("起始收入（TTM）", TTM_REV),
    ("税率", TAX),
    ("营运资金/增量收入", NWC),
    ("净负债（含融资租赁）", NET_DEBT),
    ("股权投资（非经营）", EQUITY_INV),
    ("股数（十亿）", SHARES),
]
for i, (label, val) in enumerate(inputs, start=5):
    body_cell(ws, i, 1, label, align="left")
    body_cell(ws, i, 2, val, fmt="#,##0.000" if isinstance(val, float) and abs(val) < 20 else "#,##0.0")
ws["A11"] = "注：净负债 = 35.481 + 70.973 − 64.509 = 41.945；股权投资 37.808 仅用于 SOTP，DCF/IRR 不计入。"
ws["A11"].font = Font(name=FONT, size=9, color=C_GREY)
ws.merge_cells("A11:N11")

start_rows = {"bear": 13, "base": 36, "bull": 59}
for k, start in start_rows.items():
    s = SCEN[k]
    r = RESULTS[k]
    section(ws, start, f"{s['label']}情景", 14)
    body_cell(ws, start + 1, 1, "WACC", align="left")
    body_cell(ws, start + 1, 2, s["wacc"], fmt="0.0%")
    body_cell(ws, start + 1, 3, "终值增长 g", align="left")
    body_cell(ws, start + 1, 4, s["gt"], fmt="0.0%")
    header = ["Year", "收入增速", "收入", "EBIT 利润率", "EBIT", "NOPAT",
              "D&A %", "D&A", "Capex %", "Capex", "ΔNWC", "FCFF", "折现因子", "PV"]
    header_row(ws, start + 3, header)
    for t, row in enumerate(r["rows"], start=start + 4):
        body_cell(ws, t, 1, row["year"], fmt="0")
        body_cell(ws, t, 2, s["growth"][row["year"] - 1], fmt="0.0%")
        if row["year"] == 1:
            body_cell(ws, t, 3, f"=$B$5*(1+B{t})", fmt="#,##0.0")
        else:
            body_cell(ws, t, 3, f"=C{t-1}*(1+B{t})", fmt="#,##0.0")
        body_cell(ws, t, 4, s["ebit"][row["year"] - 1], fmt="0.0%")
        body_cell(ws, t, 5, f"=C{t}*D{t}", fmt="#,##0.0")
        body_cell(ws, t, 6, f"=E{t}*(1-$B$6)", fmt="#,##0.0")
        body_cell(ws, t, 7, s["da"][row["year"] - 1], fmt="0.0%")
        body_cell(ws, t, 8, f"=C{t}*G{t}", fmt="#,##0.0")
        body_cell(ws, t, 9, s["capex"][row["year"] - 1], fmt="0.0%")
        body_cell(ws, t, 10, f"=C{t}*I{t}", fmt="#,##0.0")
        body_cell(ws, t, 11, f"=(C{t}-$B$5)*$B$7" if row["year"] == 1 else f"=(C{t}-C{t-1})*$B$7", fmt="#,##0.0")
        body_cell(ws, t, 12, f"=F{t}+H{t}-J{t}-K{t}", fmt="#,##0.0")
        body_cell(ws, t, 13, f"=1/(1+$B${start+1})^{t-start-3}", fmt="0.000")
        body_cell(ws, t, 14, f"=L{t}*M{t}", fmt="#,##0.0")
    end = start + 13
    body_cell(ws, end + 1, 1, "终值", align="left", bold=True)
    body_cell(ws, end + 1, 2, f"=L{end}*(1+$D${start+1})/($B${start+1}-$D${start+1})", fmt="#,##0.0")
    body_cell(ws, end + 2, 1, "PV 终值", align="left")
    body_cell(ws, end + 2, 2, f"=B{end+1}/(1+$B${start+1})^10", fmt="#,##0.0")
    body_cell(ws, end + 3, 1, "EV", align="left", bold=True)
    body_cell(ws, end + 3, 2, f"=SUM(N{start+4}:N{end})+B{end+2}", fmt="#,##0.0")
    body_cell(ws, end + 4, 1, "− 净负债", align="left")
    body_cell(ws, end + 4, 2, f"=$B$8", fmt="#,##0.0")
    body_cell(ws, end + 5, 1, "股权价值", align="left", bold=True)
    body_cell(ws, end + 5, 2, f"=B{end+3}-B{end+4}", fmt="#,##0.0")
    body_cell(ws, end + 6, 1, "每股价值", align="left", bold=True)
    body_cell(ws, end + 6, 2, f"=B{end+5}/$B$10", fmt="$#,##0.00")
    body_cell(ws, end + 6, 3, f"TV 占比 {r['tv_share']:.0%}；终值 FCF {r['term_fcf']:.0f}B；终值收入 {r['term_rev']:.0f}B", align="left")

chart = LineChart()
chart.title = "基准情景 FCFF 路径（十亿美元）"
chart.y_axis.title = "FCFF"
chart.x_axis.title = "Year"
data = Reference(ws, min_col=12, min_row=39, max_row=49)
cats = Reference(ws, min_col=1, min_row=40, max_row=49)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.height = 7
chart.width = 13
ws.add_chart(chart, "P13")

# --- Sheet 6: 敏感性 ---------------------------------------------------------
ws = wb.create_sheet("敏感性")
ws.sheet_view.showGridLines = False
widths(ws, {"A": 26, "B": 14, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14})
title(ws, "A1", "估值敏感性：WACC、终值 capex、终值利润率")
subtitle(ws, "A2", "基准情景；每股价值（美元）。由 Python 独立计算，Excel 表为静态值。")

header_row(ws, 4, ["WACC", "每股价值"])
for i, (w, ps) in enumerate(zip(WACC_GRID, WACC_PS), start=5):
    body_cell(ws, i, 1, w, fmt="0.0%")
    body_cell(ws, i, 2, r0(ps), fmt="$#,##0")

header_row(ws, 14, ["终值 Capex / 收入", "每股价值"])
for i, (c, ps) in enumerate(zip(CAPEX_GRID, CAPEX_PS), start=15):
    body_cell(ws, i, 1, c, fmt="0.0%")
    body_cell(ws, i, 2, r0(ps), fmt="$#,##0")

header_row(ws, 23, ["终值 EBIT 利润率", "每股价值"])
for i, (m, ps) in enumerate(zip(MARGIN_GRID, MARGIN_PS), start=24):
    body_cell(ws, i, 1, m, fmt="0.0%")
    body_cell(ws, i, 2, r0(ps), fmt="$#,##0")

section(ws, 33, "反推：当前价格隐含的 10 年收入 CAGR", 3)
body_cell(ws, 34, 1, "盘后价 $546.39", align="left")
body_cell(ws, 34, 2, IMPLIED_G_AH, fmt="0.0%")
body_cell(ws, 35, 1, "收盘价 $531.01", align="left")
body_cell(ws, 35, 2, IMPLIED_G_CLOSE, fmt="0.0%")
body_cell(ws, 36, 1, "基准情景 10 年 CAGR", align="left")
body_cell(ws, 36, 2, (RESULTS["base"]["term_rev"] / TTM_REV) ** (1 / 10) - 1, fmt="0.0%")

section(ws, 38, "SOTP 交叉检查（每股美元）", 7)
header_row(ws, 39, ["Agents & Infra EBIT 倍数", "D&C 10x", "D&C 11.5x", "D&C 13x", "", "", ""])
row = 40
for sm in [25, 28, 30, 32]:
    body_cell(ws, row, 1, f"{sm}x", align="left")
    for j, dm in enumerate([10.0, 11.5, 13.0], start=2):
        ps = next(x[3] for x in SOTP if x[0] == sm and x[1] == dm)
        body_cell(ws, row, j, r0(ps), fmt="$#,##0")
    row += 1
body_cell(ws, 45, 1, "注：A&I/DC 用 TTM EBIT 按 Q1 分部利润占比拆分；净负债扣除；股权投资按账面 37.8B 加回。", align="left")
ws.merge_cells("A45:G45")
ws["A45"].font = Font(name=FONT, size=9, color=C_GREY)

# --- Sheet 7: 来源 -----------------------------------------------------------
ws = wb.create_sheet("来源")
ws.sheet_view.showGridLines = False
widths(ws, {"A": 34, "B": 50, "C": 16, "D": 60})
title(ws, "A1", "数据来源与日期")
subtitle(ws, "A2", "研究截止 2026-10-29；价格、一致预期、目标价均为当日或前一交易日数据。")
header_row(ws, 4, ["数据", "来源", "日期", "链接/备注"])
sources = [
    ("FY2026 10-K", "SEC EDGAR, Microsoft Corp.", "2026-07-29", "https://www.sec.gov/Archives/edgar/data/789019/000119312526323660/msft-20260630.htm"),
    ("Q1 FY2027 业绩", "SEC 8-K Exhibit 99.1", "2026-10-28", "https://www.sec.gov/Archives/edgar/data/789019/000119312526436122/msft-ex99_1.htm"),
    ("Q1 FY2027 10-Q", "SEC EDGAR, Microsoft Corp.", "2026-10-28", "https://www.sec.gov/Archives/edgar/data/789019/000119312526436133/msft-20260930.htm"),
    ("Q1 FY2027 电话会", "Microsoft IR / 媒体整理", "2026-10-28", "FY2027 capex ~170B；Q2 收入 96.0–97.1B；Azure 49–50% cc；FCF 为正"),
    ("价格/市值", "Yahoo Finance / StockAnalysis / 媒体", "2026-10-28", "收盘 531.01；盘后 546.39（19:59 ET）"),
    ("一致预期", "Yahoo Finance / Investing.com / TipRanks", "2026-10", "FY2028 EPS 23.69；FY2027 EPS 19.38（Q1 前）"),
    ("目标价", "S&P Global / JPMorgan / CNBC", "2026-10-29", "均值 605.12；JPM 625；CNBC 580"),
    ("10Y 美债", "YCharts / CNBC", "2026-10-28", "5.03%"),
    ("OpenAI 股权", "TMTPost / Nasdaq / 媒体", "2025-10 至 2026-10", "27% 股权，135B 估值；2026-03 融资 852B post-money；可能稀释"),
    ("反垄断", "The Register / Samexpert / AInvest", "2026", "EU Teams 承诺具约束力；FTC 云许可调查；UK CMA"),
]
for i, (a, b, c, d) in enumerate(sources, start=5):
    body_cell(ws, i, 1, a, align="left")
    body_cell(ws, i, 2, b, align="left")
    body_cell(ws, i, 3, c, align="center")
    body_cell(ws, i, 4, d, align="left")

wb.save(XLSX)
print("saved:", XLSX)
print("scenario:", {k: round(v["ps"], 2) for k, v in RESULTS.items()})
print("IRR @546.39:", {k: round(v * 100, 2) for k, v in IRR_AH.items()})
print("IRR @531.01:", {k: round(v * 100, 2) for k, v in IRR_CLOSE.items()})
print("implied CAGR:", round(IMPLIED_G_AH * 100, 2), round(IMPLIED_G_CLOSE * 100, 2))
print("net debt:", round(NET_DEBT, 3), "EV AH:", round(EV_AH, 1))
